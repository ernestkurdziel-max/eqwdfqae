"""
Generator prezentacji:
"Opracowanie indywidualnego programu zabiegów fizykalnych
w ramach balneoklimatologii i odnowy biologicznej"

Przypadek kliniczny: Pacjent z zespołem cieśni nadgarstka.

Uruchomienie:
    pip install python-pptx
    python generuj_prezentacje.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ----------------------------------------------------------------------------
# Paleta kolorów (motyw medyczny – turkus / granat)
# ----------------------------------------------------------------------------
NAVY = RGBColor(0x0F, 0x3D, 0x5C)        # ciemny granat – tytuły
TEAL = RGBColor(0x12, 0x9A, 0x9C)        # turkus – akcent
LIGHT_TEAL = RGBColor(0xE3, 0xF4, 0xF4)  # jasny turkus – tło bloków
GRAY = RGBColor(0x44, 0x4B, 0x52)        # tekst
LIGHT_GRAY = RGBColor(0xF2, 0xF5, 0xF7)  # tło slajdu
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xF2, 0xA1, 0x3B)      # pomarańcz – wyróżnienie

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

BLANK = prs.slide_layouts[6]


# ----------------------------------------------------------------------------
# Funkcje pomocnicze
# ----------------------------------------------------------------------------
def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def add_round_rect(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def add_text(slide, x, y, w, h, text, size=18, color=GRAY, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False,
             font="Calibri", line_spacing=1.0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    if line_spacing:
        p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return tb


def add_bullets(slide, x, y, w, h, items, size=18, color=GRAY,
                bullet_color=TEAL, space_after=8, line_spacing=1.05):
    """items: lista (tekst, poziom) lub samych tekstów (poziom 0)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(space_after)
        p.line_spacing = line_spacing
        bullet = "•  " if level == 0 else "–  "
        r = p.add_run()
        r.text = bullet
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.color.rgb = bullet_color if level == 0 else ACCENT
        r.font.name = "Calibri"
        r2 = p.add_run()
        r2.text = text
        r2.font.size = Pt(size if level == 0 else size - 1)
        r2.font.color.rgb = color
        r2.font.name = "Calibri"
    return tb


def content_header(slide, number, title, subtitle=None):
    """Wspólny nagłówek slajdów treściowych."""
    set_bg(slide, LIGHT_GRAY)
    # lewy pasek akcentu
    add_rect(slide, 0, 0, Inches(0.22), SLIDE_H, TEAL)
    # numer w kółku
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.55), Inches(0.45),
                                  Inches(0.85), Inches(0.85))
    circ.fill.solid()
    circ.fill.fore_color.rgb = NAVY
    circ.line.fill.background()
    circ.shadow.inherit = False
    ctf = circ.text_frame
    ctf.word_wrap = True
    cp = ctf.paragraphs[0]
    cp.alignment = PP_ALIGN.CENTER
    cr = cp.add_run()
    cr.text = str(number)
    cr.font.size = Pt(30)
    cr.font.bold = True
    cr.font.color.rgb = WHITE
    cr.font.name = "Calibri"
    # tytuł
    add_text(slide, Inches(1.6), Inches(0.42), Inches(11.2), Inches(0.7),
             title, size=30, color=NAVY, bold=True)
    if subtitle:
        add_text(slide, Inches(1.62), Inches(1.05), Inches(11.2), Inches(0.4),
                 subtitle, size=15, color=TEAL, italic=True)
    # linia oddzielająca
    add_rect(slide, Inches(1.6), Inches(1.5), Inches(11.1), Pt(2), TEAL)


# ============================================================================
# SLAJD 1 – Tytułowy
# ============================================================================
s = add_slide()
set_bg(s, NAVY)
# pas akcentu
add_rect(s, 0, Inches(2.7), SLIDE_W, Inches(0.08), TEAL)
add_rect(s, 0, Inches(5.05), SLIDE_W, Inches(0.04), ACCENT)

add_text(s, Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.5),
         "BALNEOKLIMATOLOGIA I ODNOWA BIOLOGICZNA", size=16, color=TEAL,
         bold=True, align=PP_ALIGN.CENTER)

add_text(s, Inches(0.8), Inches(2.85), Inches(11.7), Inches(1.6),
         "Indywidualny program zabiegów fizykalnych", size=40, color=WHITE,
         bold=True, align=PP_ALIGN.CENTER)

add_text(s, Inches(0.8), Inches(4.15), Inches(11.7), Inches(0.8),
         "Przypadek kliniczny: pacjent z zespołem cieśni nadgarstka",
         size=22, color=ACCENT, bold=True, align=PP_ALIGN.CENTER)

add_text(s, Inches(0.8), Inches(5.3), Inches(11.7), Inches(0.6),
         "Opracowanie indywidualnego programu w ramach balneoterapii "
         "i odnowy biologicznej", size=15, color=LIGHT_TEAL,
         align=PP_ALIGN.CENTER, italic=True)

add_text(s, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.5),
         "Imię i nazwisko  •  kierunek studiów  •  rok akademicki",
         size=13, color=RGBColor(0xA9, 0xC4, 0xD4), align=PP_ALIGN.CENTER)


# ============================================================================
# SLAJD 2 – Plan / spis treści
# ============================================================================
s = add_slide()
content_header(s, "i", "Plan prezentacji")

left_items = [
    "1.  Charakterystyka pacjenta",
    "2.  Analiza wskazań i przeciwwskazań",
    "3.  Dobór zabiegów fizykalnych",
    "4.  Opracowanie programu terapii",
]
right_items = [
    "5.  Efekty terapeutyczne",
    "6.  Elementy odnowy biologicznej",
    "7.  Podsumowanie i wnioski",
    "8.  Bibliografia",
]
y0 = Inches(2.1)
for i, txt in enumerate(left_items):
    box = add_round_rect(s, Inches(1.6), y0 + Inches(i * 1.15), Inches(5.2),
                         Inches(0.9), WHITE)
    add_text(s, Inches(1.8), y0 + Inches(i * 1.15) + Inches(0.18), Inches(4.9),
             Inches(0.6), txt, size=17, color=NAVY, bold=True)
for i, txt in enumerate(right_items):
    box = add_round_rect(s, Inches(7.15), y0 + Inches(i * 1.15), Inches(5.2),
                         Inches(0.9), WHITE)
    add_text(s, Inches(7.35), y0 + Inches(i * 1.15) + Inches(0.18),
             Inches(4.9), Inches(0.6), txt, size=17, color=NAVY, bold=True)


# ============================================================================
# SLAJD 3 – Charakterystyka pacjenta
# ============================================================================
s = add_slide()
content_header(s, 1, "Charakterystyka pacjenta",
               "Dane podstawowe, tryb życia i rozpoznanie")

# karta danych po lewej
card = add_round_rect(s, Inches(1.6), Inches(1.8), Inches(4.6), Inches(4.9),
                      WHITE)
add_text(s, Inches(1.85), Inches(2.0), Inches(4.2), Inches(0.5),
         "Dane podstawowe", size=18, color=TEAL, bold=True)
data = [
    ("Wiek:", "39 lat"),
    ("Płeć:", "mężczyzna"),
    ("Zawód:", "informatyk"),
    ("Praca:", "ponad 8 h dziennie przy komputerze"),
    ("Tryb życia:", "siedzący, niska aktywność fizyczna"),
]
yy = 2.55
for label, val in data:
    add_text(s, Inches(1.85), Inches(yy), Inches(4.1), Inches(0.4),
             label, size=15, color=NAVY, bold=True)
    add_text(s, Inches(1.85), Inches(yy + 0.32), Inches(4.1), Inches(0.5),
             val, size=14, color=GRAY)
    yy += 0.82

# rozpoznanie i dolegliwości po prawej
add_text(s, Inches(6.6), Inches(1.85), Inches(6.1), Inches(0.5),
         "Rozpoznanie", size=18, color=TEAL, bold=True)
add_bullets(s, Inches(6.6), Inches(2.3), Inches(6.1), Inches(1.0), [
    "Zespół cieśni nadgarstka (ucisk nerwu pośrodkowego "
    "w kanale nadgarstka) prawej ręki.",
], size=15)

add_text(s, Inches(6.6), Inches(3.35), Inches(6.1), Inches(0.5),
         "Dolegliwości i ograniczenia funkcjonalne", size=18, color=TEAL,
         bold=True)
add_bullets(s, Inches(6.6), Inches(3.85), Inches(6.1), Inches(2.7), [
    "Drętwienie i mrowienie palców prawej dłoni.",
    "Nasilenie objawów w nocy oraz podczas pracy przy klawiaturze.",
    "Ból nadgarstka, osłabienie siły chwytu.",
    "Trudności w wykonywaniu precyzyjnych czynności manualnych.",
    "Przeciążenie wynikające z monotonnej, statycznej pozycji.",
], size=15, space_after=9)


# ============================================================================
# SLAJD 4 – Analiza wskazań
# ============================================================================
s = add_slide()
content_header(s, 2, "Analiza wskazań",
               "Wskazania do zastosowania zabiegów fizykalnych")

add_bullets(s, Inches(1.7), Inches(1.85), Inches(11.0), Inches(3.0), [
    "Zespół cieśni nadgarstka o łagodnym / umiarkowanym nasileniu "
    "(bez wskazań do leczenia operacyjnego).",
    "Dolegliwości bólowe i parestezje (drętwienie, mrowienie).",
    "Stan przeciążeniowy ścięgien i pochewek w okolicy nadgarstka.",
    "Osłabienie siły mięśniowej i obniżona sprawność manualna ręki.",
    "Przewlekłe napięcie mięśniowe obręczy barkowej i przedramienia "
    "wynikające z pracy siedzącej.",
    "Profilaktyka wtórna – zapobieganie nawrotom i progresji objawów.",
], size=18, space_after=12)

# uzasadnienie – pasek na dole
add_round_rect(s, Inches(1.7), Inches(5.75), Inches(11.0), Inches(1.2),
               LIGHT_TEAL)
add_text(s, Inches(1.95), Inches(5.9), Inches(10.6), Inches(0.4),
         "Uzasadnienie wyboru metod", size=15, color=TEAL, bold=True)
add_text(s, Inches(1.95), Inches(6.28), Inches(10.6), Inches(0.6),
         "Metody fizykalne działają przeciwbólowo, przeciwzapalnie i "
         "przeciwobrzękowo, poprawiają ukrwienie oraz zmniejszają ucisk na "
         "nerw pośrodkowy – są leczeniem zachowawczym pierwszego wyboru.",
         size=13, color=GRAY, line_spacing=1.05)


# ============================================================================
# SLAJD 5 – Przeciwwskazania
# ============================================================================
s = add_slide()
content_header(s, 2, "Przeciwwskazania",
               "Przeciwwskazania ogólne i miejscowe")

# dwie kolumny
add_round_rect(s, Inches(1.6), Inches(1.85), Inches(5.4), Inches(4.85), WHITE)
add_rect(s, Inches(1.6), Inches(1.85), Inches(5.4), Inches(0.6), NAVY)
add_text(s, Inches(1.8), Inches(1.95), Inches(5.0), Inches(0.4),
         "Przeciwwskazania ogólne", size=16, color=WHITE, bold=True)
add_bullets(s, Inches(1.85), Inches(2.65), Inches(5.0), Inches(4.0), [
    "Choroba nowotworowa (czynna).",
    "Stany gorączkowe i ostre infekcje.",
    "Niewydolność krążenia i ciężkie choroby serca.",
    "Skłonność do krwawień, leczenie przeciwkrzepliwe.",
    "Ciąża (dla wybranych zabiegów).",
    "Rozrusznik serca (zabiegi elektro- i magnetoterapii).",
], size=14, space_after=9)

add_round_rect(s, Inches(7.3), Inches(1.85), Inches(5.4), Inches(4.85), WHITE)
add_rect(s, Inches(7.3), Inches(1.85), Inches(5.4), Inches(0.6), TEAL)
add_text(s, Inches(7.5), Inches(1.95), Inches(5.0), Inches(0.4),
         "Przeciwwskazania miejscowe", size=16, color=WHITE, bold=True)
add_bullets(s, Inches(7.55), Inches(2.65), Inches(5.0), Inches(4.0), [
    "Ostry stan zapalny i obrzęk w okolicy nadgarstka.",
    "Uszkodzenia i zmiany skórne, rany, infekcje skóry.",
    "Świeży uraz, złamanie w obrębie ręki.",
    "Zaawansowany zanik mięśni kłębu kciuka "
    "(wskazanie do konsultacji chirurgicznej).",
    "Zakrzepica żył kończyny.",
], size=14, space_after=9)


# ============================================================================
# SLAJD 6 – Dobór zabiegów (część 1: balneoterapia / hydroterapia)
# ============================================================================
s = add_slide()
content_header(s, 3, "Dobór zabiegów fizykalnych (1/2)",
               "Balneoterapia, hydroterapia i zabiegi cieplne")

rows = [
    ("Kąpiele solankowe / borowinowe ręki",
     "Działanie przeciwbólowe, przeciwzapalne, poprawa ukrwienia tkanek.",
     "Zmniejszenie obrzęku i ucisku na nerw pośrodkowy."),
    ("Kąpiel wirowa kończyny górnej",
     "Delikatny masaż wodny + ciepło – rozluźnienie i lepsze krążenie.",
     "Redukcja napięcia mięśni przedramienia, poprawa trofiki."),
    ("Okłady borowinowe (zawijania)",
     "Głębokie, długotrwałe przegrzanie + składniki biologicznie czynne.",
     "Działanie przeciwzapalne i regenerujące tkanki okołostawowe."),
    ("Kąpiele naprzemienne (cieplno-zimne)",
     "Naprzemienne rozszerzanie i obkurczanie naczyń – „gimnastyka naczyń”.",
     "Poprawa mikrokrążenia, zmniejszenie zastoju i obrzęku."),
]
y = 1.85
add_rect(s, Inches(1.6), Inches(y), Inches(4.0), Inches(0.5), NAVY)
add_rect(s, Inches(5.6), Inches(y), Inches(4.0), Inches(0.5), NAVY)
add_rect(s, Inches(9.6), Inches(y), Inches(3.1), Inches(0.5), NAVY)
add_text(s, Inches(1.7), Inches(y + 0.06), Inches(3.8), Inches(0.4),
         "Zabieg", size=13, color=WHITE, bold=True)
add_text(s, Inches(5.7), Inches(y + 0.06), Inches(3.8), Inches(0.4),
         "Mechanizm działania", size=13, color=WHITE, bold=True)
add_text(s, Inches(9.7), Inches(y + 0.06), Inches(2.9), Inches(0.4),
         "Cel terapeutyczny", size=13, color=WHITE, bold=True)
y += 0.5
rh = 1.18
for i, (z, m, c) in enumerate(rows):
    bg = WHITE if i % 2 == 0 else LIGHT_TEAL
    add_rect(s, Inches(1.6), Inches(y), Inches(11.1), Inches(rh), bg)
    add_text(s, Inches(1.7), Inches(y + 0.08), Inches(3.8), Inches(rh - 0.1),
             z, size=12.5, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(5.7), Inches(y + 0.08), Inches(3.8), Inches(rh - 0.1),
             m, size=11.5, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(9.7), Inches(y + 0.08), Inches(2.9), Inches(rh - 0.1),
             c, size=11.5, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
    y += rh


# ============================================================================
# SLAJD 7 – Dobór zabiegów (część 2: fizykoterapia + odnowa)
# ============================================================================
s = add_slide()
content_header(s, 3, "Dobór zabiegów fizykalnych (2/2)",
               "Fizykoterapia, masaż i kinezyterapia")

rows = [
    ("Krioterapia miejscowa",
     "Krótkotrwałe schłodzenie – efekt przeciwbólowy i przeciwobrzękowy.",
     "Ograniczenie stanu zapalnego pochewek ścięgnistych."),
    ("Ultradźwięki / fonoforeza",
     "Mikromasaż tkankowy, działanie przeciwbólowe i przeciwzapalne.",
     "Zmniejszenie dolegliwości w obrębie kanału nadgarstka."),
    ("Laseroterapia / magnetoterapia",
     "Stymulacja regeneracji nerwu, działanie przeciwzapalne.",
     "Poprawa przewodnictwa nerwowego, redukcja parestezji."),
    ("Masaż przedramienia i obręczy barkowej",
     "Rozluźnienie mięśni, poprawa krążenia i drenażu.",
     "Zmniejszenie napięcia i przeciążeń wtórnych."),
    ("Kinezyterapia – ćwiczenia ślizgowe nerwu",
     "Mobilizacja nerwu i ścięgien, ćwiczenia wzmacniające.",
     "Poprawa ruchomości, siły chwytu i sprawności ręki."),
]
y = 1.85
add_rect(s, Inches(1.6), Inches(y), Inches(4.0), Inches(0.5), TEAL)
add_rect(s, Inches(5.6), Inches(y), Inches(4.0), Inches(0.5), TEAL)
add_rect(s, Inches(9.6), Inches(y), Inches(3.1), Inches(0.5), TEAL)
add_text(s, Inches(1.7), Inches(y + 0.06), Inches(3.8), Inches(0.4),
         "Zabieg", size=13, color=WHITE, bold=True)
add_text(s, Inches(5.7), Inches(y + 0.06), Inches(3.8), Inches(0.4),
         "Mechanizm działania", size=13, color=WHITE, bold=True)
add_text(s, Inches(9.7), Inches(y + 0.06), Inches(2.9), Inches(0.4),
         "Cel terapeutyczny", size=13, color=WHITE, bold=True)
y += 0.5
rh = 0.95
for i, (z, m, c) in enumerate(rows):
    bg = WHITE if i % 2 == 0 else LIGHT_TEAL
    add_rect(s, Inches(1.6), Inches(y), Inches(11.1), Inches(rh), bg)
    add_text(s, Inches(1.7), Inches(y + 0.05), Inches(3.8), Inches(rh - 0.05),
             z, size=12, color=NAVY, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(5.7), Inches(y + 0.05), Inches(3.8), Inches(rh - 0.05),
             m, size=11, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(9.7), Inches(y + 0.05), Inches(2.9), Inches(rh - 0.05),
             c, size=11, color=GRAY, anchor=MSO_ANCHOR.MIDDLE)
    y += rh


# ============================================================================
# SLAJD 8 – Program terapii (plan tygodniowy)
# ============================================================================
s = add_slide()
content_header(s, 4, "Opracowanie programu terapii",
               "Przykładowy plan tygodniowy (cykl 3 tygodni / 15 zabiegów)")

days = [
    ("Pon.", "Kąpiel wirowa ręki (15 min)\nUltradźwięki (5 min)\nĆwiczenia ślizgowe nerwu"),
    ("Wt.", "Okład borowinowy (20 min)\nMagnetoterapia (15 min)"),
    ("Śr.", "Kąpiel solankowa (15 min)\nMasaż przedramienia (15 min)"),
    ("Czw.", "Laseroterapia\nKinezyterapia – wzmacnianie\nĆwiczenia rozciągające"),
    ("Pt.", "Kąpiele naprzemienne\nKrioterapia (w razie bólu)\nRelaksacja"),
]
n = len(days)
total_w = 11.1
gap = 0.2
cw = (total_w - gap * (n - 1)) / n
x = 1.6
y = 2.0
for d, content in days:
    add_rect(s, Inches(x), Inches(y), Inches(cw), Inches(0.55), NAVY)
    add_text(s, Inches(x), Inches(y + 0.08), Inches(cw), Inches(0.4),
             d, size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_round_rect(s, Inches(x), Inches(y + 0.62), Inches(cw), Inches(2.9),
                   WHITE)
    tb = s.shapes.add_textbox(Inches(x + 0.05), Inches(y + 0.72),
                              Inches(cw - 0.1), Inches(2.7))
    tf = tb.text_frame
    tf.word_wrap = True
    for j, line in enumerate(content.split("\n")):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_after = Pt(6)
        r = p.add_run()
        r.text = line
        r.font.size = Pt(10.5)
        r.font.color.rgb = GRAY
        r.font.name = "Calibri"
    x += cw + gap

add_round_rect(s, Inches(1.6), Inches(5.7), Inches(11.1), Inches(1.25),
               LIGHT_TEAL)
add_text(s, Inches(1.85), Inches(5.82), Inches(10.6), Inches(0.4),
         "Zasady realizacji programu", size=14, color=TEAL, bold=True)
add_bullets(s, Inches(1.85), Inches(6.2), Inches(10.6), Inches(0.7), [
    "Częstotliwość: 5 dni w tygodniu; weekend – odpoczynek i autoterapia "
    "(ćwiczenia, higiena pracy). Czas serii: ok. 3 tygodnie (15 zabiegów).",
    "Kolejność: zabiegi cieplne / rozluźniające → terapia właściwa → "
    "kinezyterapia. Intensywność dostosowana do reakcji i stanu pacjenta.",
], size=11.5, space_after=4)


# ============================================================================
# SLAJD 9 – Efekty terapeutyczne
# ============================================================================
s = add_slide()
content_header(s, 5, "Efekty terapeutyczne",
               "Przewidywane rezultaty zastosowanej terapii")

cards = [
    ("Zmniejszenie bólu", "Redukcja dolegliwości bólowych nadgarstka "
     "i parestezji (drętwienia, mrowienia) palców."),
    ("Lepsze ukrwienie", "Poprawa mikrokrążenia i trofiki tkanek, "
     "zmniejszenie obrzęku i ucisku na nerw pośrodkowy."),
    ("Większa sprawność", "Wzrost siły chwytu i precyzji ruchów, "
     "poprawa zakresu ruchomości nadgarstka."),
    ("Rozluźnienie", "Zmniejszenie napięcia mięśni przedramienia "
     "i obręczy barkowej, ustąpienie przeciążeń."),
    ("Lepszy sen", "Ograniczenie nocnych dolegliwości bólowych "
     "poprawia jakość snu i regenerację."),
    ("Lepsze samopoczucie", "Ogólna regeneracja, redukcja stresu "
     "i poprawa komfortu pracy oraz życia."),
]
cw, ch = 3.55, 1.95
gx, gy = 0.25, 0.25
x0, y0 = 1.7, 2.0
for i, (t, d) in enumerate(cards):
    col = i % 3
    row = i // 3
    x = x0 + col * (cw + gx)
    y = y0 + row * (ch + gy)
    add_round_rect(s, Inches(x), Inches(y), Inches(cw), Inches(ch), WHITE)
    add_rect(s, Inches(x), Inches(y), Inches(0.12), Inches(ch), ACCENT)
    add_text(s, Inches(x + 0.3), Inches(y + 0.15), Inches(cw - 0.4),
             Inches(0.5), t, size=16, color=NAVY, bold=True)
    add_text(s, Inches(x + 0.3), Inches(y + 0.65), Inches(cw - 0.45),
             Inches(1.2), d, size=12, color=GRAY, line_spacing=1.05)


# ============================================================================
# SLAJD 10 – Elementy odnowy biologicznej
# ============================================================================
s = add_slide()
content_header(s, 6, "Elementy odnowy biologicznej",
               "Dodatkowe formy regeneracji i ich znaczenie")

items = [
    ("Ergonomia i higiena pracy",
     "Prawidłowe ustawienie stanowiska, podkładka pod nadgarstek, "
     "regularne przerwy i mikroćwiczenia co 45–60 min."),
    ("Aktywność fizyczna",
     "Regularny ruch (pływanie, spacery), ćwiczenia ogólnousprawniające – "
     "przeciwdziałanie skutkom siedzącego trybu życia."),
    ("Techniki relaksacyjne",
     "Trening relaksacyjny, ćwiczenia oddechowe i rozciągające – "
     "redukcja napięcia mięśniowego i stresu."),
    ("Dieta i nawodnienie",
     "Dieta przeciwzapalna, bogata w wit. z grupy B; odpowiednie "
     "nawodnienie wspiera regenerację tkanki nerwowej."),
    ("Higiena snu i odpoczynek",
     "Regularny sen, unikanie spania z mocno zgiętym nadgarstkiem, "
     "orteza nocna stabilizująca rękę."),
    ("Edukacja zdrowotna",
     "Świadomość czynników ryzyka i profilaktyka nawrotów – "
     "warunek trwałości efektów terapii."),
]
cw, ch = 5.45, 1.45
gx, gy = 0.2, 0.2
x0, y0 = 1.7, 1.95
for i, (t, d) in enumerate(items):
    col = i % 2
    row = i // 2
    x = x0 + col * (cw + gx)
    y = y0 + row * (ch + gy)
    add_round_rect(s, Inches(x), Inches(y), Inches(cw), Inches(ch), WHITE)
    num = slide_num = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.18),
                                         Inches(y + 0.28), Inches(0.55),
                                         Inches(0.55))
    num.fill.solid()
    num.fill.fore_color.rgb = TEAL
    num.line.fill.background()
    num.shadow.inherit = False
    ntf = num.text_frame
    np = ntf.paragraphs[0]
    np.alignment = PP_ALIGN.CENTER
    nr = np.add_run()
    nr.text = str(i + 1)
    nr.font.size = Pt(18)
    nr.font.bold = True
    nr.font.color.rgb = WHITE
    add_text(s, Inches(x + 0.95), Inches(y + 0.12), Inches(cw - 1.1),
             Inches(0.45), t, size=14.5, color=NAVY, bold=True)
    add_text(s, Inches(x + 0.95), Inches(y + 0.55), Inches(cw - 1.1),
             Inches(0.85), d, size=11.5, color=GRAY, line_spacing=1.0)


# ============================================================================
# SLAJD 11 – Podsumowanie i wnioski
# ============================================================================
s = add_slide()
content_header(s, 7, "Podsumowanie i wnioski",
               "Ocena skuteczności i możliwości modyfikacji terapii")

add_text(s, Inches(1.7), Inches(1.85), Inches(11.0), Inches(0.5),
         "Ocena skuteczności programu", size=18, color=TEAL, bold=True)
add_bullets(s, Inches(1.7), Inches(2.35), Inches(11.0), Inches(2.0), [
    "Kompleksowe leczenie zachowawcze (balneoterapia + fizykoterapia + "
    "kinezyterapia + odnowa biologiczna) skutecznie zmniejsza objawy "
    "łagodnego i umiarkowanego zespołu cieśni nadgarstka.",
    "Połączenie zabiegów miejscowych z modyfikacją stylu życia i ergonomii "
    "pracy daje trwalsze efekty niż samo leczenie objawowe.",
    "Wczesne wdrożenie terapii zmniejsza ryzyko progresji i potrzeby "
    "leczenia operacyjnego.",
], size=15, space_after=9)

add_text(s, Inches(1.7), Inches(5.0), Inches(11.0), Inches(0.5),
         "Możliwości modyfikacji terapii", size=18, color=TEAL, bold=True)
add_bullets(s, Inches(1.7), Inches(5.5), Inches(11.0), Inches(1.6), [
    "Program należy monitorować i dostosowywać do reakcji oraz postępów "
    "pacjenta (intensywność, dobór i liczba zabiegów).",
    "Brak poprawy lub nasilenie objawów / zanik mięśni → konsultacja "
    "neurologiczna i chirurgiczna (EMG, ewentualne leczenie operacyjne).",
], size=15, space_after=9)


# ============================================================================
# SLAJD 12 – Bibliografia
# ============================================================================
s = add_slide()
content_header(s, 8, "Bibliografia",
               "Literatura podstawowa i uzupełniająca")

add_text(s, Inches(1.7), Inches(1.8), Inches(11.0), Inches(0.4),
         "Literatura podstawowa", size=16, color=TEAL, bold=True)
add_bullets(s, Inches(1.7), Inches(2.25), Inches(11.0), Inches(2.4), [
    "Straburzyński G., Straburzyńska-Lupa A.: Medycyna fizykalna. "
    "PZWL, Warszawa 2000.",
    "Ponikowska I., Ferson D.: Nowoczesna Medycyna Uzdrowiskowa. "
    "Medi Press, Warszawa 2009.",
    "Kasprzak W.: Fizjoterapia kliniczna. PZWL, Warszawa 2011.",
    "Kasprzak W., Mańkowska A.: Fizykoterapia, medycyna uzdrowiskowa "
    "i SPA. PZWL, Warszawa 2010.",
], size=14, space_after=8)

add_text(s, Inches(1.7), Inches(4.75), Inches(11.0), Inches(0.4),
         "Literatura uzupełniająca", size=16, color=TEAL, bold=True)
add_bullets(s, Inches(1.7), Inches(5.2), Inches(11.0), Inches(1.5), [
    "Ponikowska J.: Kompendium balneologii: rekomendacje krajowego "
    "konsultanta. Toruń 2012.",
    "Kochański J. W.: Balneologia i hydroterapia. Wydawnictwo AWF, "
    "Wrocław 2002.",
], size=14, space_after=8)


# ============================================================================
# SLAJD 13 – Dziękuję
# ============================================================================
s = add_slide()
set_bg(s, NAVY)
add_rect(s, 0, Inches(3.35), SLIDE_W, Inches(0.06), TEAL)
add_text(s, Inches(0.8), Inches(2.7), Inches(11.7), Inches(1.0),
         "Dziękuję za uwagę", size=40, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Inches(0.8), Inches(3.7), Inches(11.7), Inches(0.6),
         "Indywidualny program zabiegów fizykalnych – "
         "zespół cieśni nadgarstka", size=16, color=LIGHT_TEAL,
         align=PP_ALIGN.CENTER, italic=True)


# ----------------------------------------------------------------------------
prs.save("Zespol_ciesni_nadgarstka_prezentacja.pptx")
print("Zapisano: Zespol_ciesni_nadgarstka_prezentacja.pptx")
print(f"Liczba slajdów: {len(prs.slides._sldIdLst)}")
