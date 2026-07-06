from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTFILE = "Mukowiscydoza_fizjoterapia_Ernest_Kurdziel_72714.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

W = prs.slide_width
H = prs.slide_height

NAVY = RGBColor(22, 49, 83)
TEAL = RGBColor(0, 125, 133)
CYAN = RGBColor(74, 170, 185)
LIGHT = RGBColor(238, 247, 249)
PALE = RGBColor(225, 239, 242)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(38, 45, 52)
MUTED = RGBColor(92, 105, 116)
ORANGE = RGBColor(232, 134, 58)
GREEN = RGBColor(82, 155, 105)


def add_bg(slide, section=None):
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = WHITE
    band = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, Inches(0.28)
    )
    band.fill.solid()
    band.fill.fore_color.rgb = TEAL
    band.line.fill.background()
    if section:
        tag = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(10.65),
            Inches(0.45),
            Inches(2.15),
            Inches(0.38),
        )
        tag.fill.solid()
        tag.fill.fore_color.rgb = LIGHT
        tag.line.color.rgb = PALE
        tf = tag.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = section
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = TEAL


def add_title(slide, title, subtitle=None, section=None):
    add_bg(slide, section)
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.5), Inches(9.8), Inches(0.7))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(28)
    r.font.bold = True
    r.font.color.rgb = NAVY
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.58), Inches(1.16), Inches(10.6), Inches(0.35))
        stf = sub.text_frame
        stf.clear()
        p = stf.paragraphs[0]
        r = p.add_run()
        r.text = subtitle
        r.font.size = Pt(12)
        r.font.color.rgb = MUTED


def add_footer(slide, n, citation=None):
    line = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.55), Inches(7.05), Inches(12.25), Inches(0.01)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = PALE
    line.line.fill.background()
    txt = f"{n}/30  |  Ernest Kurdziel, nr albumu 72714"
    if citation:
        txt += f"  |  {citation}"
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.1), Inches(12.25), Inches(0.25))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = txt
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED


def bullet_box(slide, items, x=0.8, y=1.65, w=7.1, h=4.9, font_size=18, color=DARK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        if isinstance(item, tuple):
            text, level = item
        else:
            text, level = item, 0
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.level = level
        p.font.size = Pt(font_size if level == 0 else font_size - 2)
        p.font.color.rgb = color
        p.space_after = Pt(7)
    return box


def two_col(slide, left_title, left_items, right_title, right_items, y=1.55):
    for x, title, items, c in [
        (0.7, left_title, left_items, TEAL),
        (6.85, right_title, right_items, ORANGE),
    ]:
        header = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y),
            Inches(5.75),
            Inches(0.55),
        )
        header.fill.solid()
        header.fill.fore_color.rgb = c
        header.line.fill.background()
        tf = header.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = title
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = WHITE
        panel = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(x),
            Inches(y + 0.72),
            Inches(5.75),
            Inches(4.75),
        )
        panel.fill.solid()
        panel.fill.fore_color.rgb = LIGHT if c == TEAL else RGBColor(255, 244, 235)
        panel.line.color.rgb = PALE
        bullet_box(slide, items, x + 0.28, y + 1.0, 5.15, 4.2, 15)


def callout(slide, text, x, y, w, h, fill=LIGHT, line=CYAN, font=15):
    shape = slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.14)
    tf.margin_right = Inches(0.14)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.size = Pt(font)
    r.font.bold = True
    r.font.color.rgb = NAVY
    return shape


def add_cycle(slide, labels, center=(6.65, 3.75), radius=2.25):
    colors = [TEAL, CYAN, GREEN, ORANGE, RGBColor(117, 100, 170), RGBColor(47, 121, 180)]
    positions = [
        (center[0], center[1] - radius),
        (center[0] + radius * 0.9, center[1] - radius * 0.35),
        (center[0] + radius * 0.56, center[1] + radius * 0.75),
        (center[0] - radius * 0.56, center[1] + radius * 0.75),
        (center[0] - radius * 0.9, center[1] - radius * 0.35),
    ]
    for i, lab in enumerate(labels):
        x, y = positions[i]
        callout(slide, lab, x - 1.05, y - 0.32, 2.1, 0.64, RGBColor(245, 250, 251), colors[i], 12)
    callout(slide, "indywidualizacja\n+ monitorowanie", center[0] - 1.05, center[1] - 0.47, 2.1, 0.94, PALE, TEAL, 12)


slides = []

# 1
s = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(s)
cover = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, W, H)
cover.fill.solid()
cover.fill.fore_color.rgb = NAVY
cover.line.fill.background()
accent = s.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(0.38), H)
accent.fill.solid()
accent.fill.fore_color.rgb = TEAL
accent.line.fill.background()
title = s.shapes.add_textbox(Inches(0.85), Inches(1.05), Inches(11.3), Inches(1.55))
tf = title.text_frame
tf.clear()
p = tf.paragraphs[0]
r = p.add_run()
r.text = "Zasady planowania i programowania fizjoterapii pacjentów z chorobami układu oddechowego"
r.font.size = Pt(30)
r.font.bold = True
r.font.color.rgb = WHITE
sub = s.shapes.add_textbox(Inches(0.9), Inches(3.05), Inches(10.8), Inches(1.2))
tf = sub.text_frame
tf.clear()
p = tf.paragraphs[0]
r = p.add_run()
r.text = "Jednostka chorobowa: mukowiscydoza"
r.font.size = Pt(26)
r.font.bold = True
r.font.color.rgb = CYAN
p = tf.add_paragraph()
p.text = "Przedmiot: kliniczne podstawy fizjoterapii w pulmonologii"
p.font.size = Pt(17)
p.font.color.rgb = WHITE
meta = s.shapes.add_textbox(Inches(0.9), Inches(5.4), Inches(8.2), Inches(0.8))
tf = meta.text_frame
tf.clear()
p = tf.paragraphs[0]
p.text = "Autor: Ernest Kurdziel  |  Numer albumu: 72714"
p.font.size = Pt(17)
p.font.color.rgb = WHITE
add_footer(s, 1)

# 2
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Cel i zakres prezentacji", section="wprowadzenie")
bullet_box(
    s,
    [
        "przedstawienie mukowiscydozy jako choroby układu oddechowego wymagającej stałej fizjoterapii",
        "omówienie badania fizjoterapeutycznego pacjenta: od wywiadu po testy funkcjonalne",
        "dobór technik oczyszczania dróg oddechowych, treningu i edukacji",
        "zaproponowanie planu terapii oraz kryteriów monitorowania efektów",
        "oparcie zaleceń na literaturze z PubMed / Google Scholar",
    ],
    x=1.0,
    y=1.65,
    w=10.9,
    h=4.7,
    font_size=18,
)
add_footer(s, 2)

# 3
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Mukowiscydoza: definicja i istota problemu", section="choroba")
two_col(
    s,
    "Czym jest CF?",
    [
        "choroba genetyczna związana z mutacjami genu CFTR",
        "zaburzenie transportu jonów chlorkowych i wody",
        "gęsta, lepka wydzielina w drogach oddechowych",
        "przewlekłe zakażenia, stan zapalny i postępujące uszkodzenie płuc",
    ],
    "Znaczenie dla fizjoterapii",
    [
        "utrudniony transport śluzowo-rzęskowy",
        "nawracające zaostrzenia oskrzelowo-płucne",
        "spadek tolerancji wysiłku i duszność",
        "codzienna terapia jest częścią leczenia podtrzymującego",
    ],
)
add_footer(s, 3, "Flume et al., 2009")

# 4
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Patofizjologia zmian w układzie oddechowym", section="choroba")
callout(s, "mutacja CFTR", 0.8, 1.7, 2.0, 0.7, PALE, TEAL)
callout(s, "odwodnienie\npowierzchni dróg oddechowych", 3.25, 1.7, 2.4, 0.8, LIGHT, CYAN, 13)
callout(s, "zaleganie\nwydzieliny", 6.1, 1.7, 2.0, 0.7, PALE, TEAL)
callout(s, "infekcja +\nstan zapalny", 8.55, 1.7, 2.0, 0.7, RGBColor(255, 244, 235), ORANGE)
callout(s, "obturacja,\nrozstrzenie, spadek FEV1", 10.95, 1.65, 1.75, 0.9, RGBColor(255, 244, 235), ORANGE, 12)
bullet_box(
    s,
    [
        "wydzielina tworzy środowisko sprzyjające kolonizacji bakteryjnej",
        "obturacja nasila nierównomierną wentylację i pracę oddychania",
        "przewlekły stan zapalny prowadzi do przebudowy oskrzeli",
        "fizjoterapia celuje w transport wydzieliny, wentylację, tolerancję wysiłku i samodzielność pacjenta",
    ],
    x=1.05,
    y=3.25,
    w=11.0,
    h=3.2,
    font_size=18,
)
add_footer(s, 4, "Flume et al., 2009")

# 5
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Obraz kliniczny istotny dla fizjoterapeuty", section="choroba")
two_col(
    s,
    "Objawy oddechowe",
    [
        "przewlekły kaszel, często produktywny",
        "nawracające infekcje i zaostrzenia",
        "świsty, duszność, pogorszenie tolerancji wysiłku",
        "krwioplucie lub ból w klatce piersiowej jako sygnały alarmowe",
    ],
    "Objawy ogólnoustrojowe",
    [
        "niedożywienie lub trudności w utrzymaniu masy ciała",
        "niewydolność zewnątrzwydzielnicza trzustki",
        "cukrzyca związana z CF, osteopenia, zmęczenie",
        "obciążenie psychiczne i wysoki ciężar codziennego leczenia",
    ],
)
add_footer(s, 5, "Castellani et al., 2018")

# 6
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Nadrzędne cele fizjoterapii w mukowiscydozie", section="cele terapii")
for x, y, text, col in [
    (0.9, 1.55, "utrzymanie drożności dróg oddechowych", TEAL),
    (4.7, 1.55, "zmniejszenie zalegania wydzieliny", CYAN),
    (8.5, 1.55, "profilaktyka zaostrzeń", GREEN),
    (0.9, 3.15, "poprawa wydolności i siły", ORANGE),
    (4.7, 3.15, "utrzymanie ruchomości klatki piersiowej", RGBColor(117, 100, 170)),
    (8.5, 3.15, "samodzielność i adherencja", RGBColor(47, 121, 180)),
]:
    callout(s, text, x, y, 3.25, 1.05, RGBColor(245, 250, 251), col, 15)
bullet_box(
    s,
    ["Cele powinny być mierzalne, uzgodnione z pacjentem i modyfikowane przy zmianie stanu klinicznego."],
    x=1.25,
    y=5.1,
    w=10.8,
    h=0.9,
    font_size=19,
)
add_footer(s, 6, "McIlwaine et al., 2019")

# 7
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Zasady planowania programu terapii", section="planowanie")
add_cycle(s, ["badanie", "cele SMART", "dobór technik", "edukacja", "kontrola efektów"])
bullet_box(
    s,
    [
        "program nie jest stały: zmienia się z wiekiem, zaostrzeniami, preferencjami i wynikami badań",
        "dobór ACT uwzględnia skuteczność, bezpieczeństwo, czas, sprzęt i akceptację pacjenta",
        "ćwiczenia są uzupełnieniem fizjoterapii oddechowej, a nie jedynym substytutem oczyszczania",
        "najważniejszy wynik praktyczny: regularność możliwa do utrzymania w codziennym życiu",
    ],
    x=0.8,
    y=1.65,
    w=4.7,
    h=4.9,
    font_size=15,
)
add_footer(s, 7, "Flume et al., 2009; Williams et al., 2022")

# 8
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Badanie pacjenta: schemat całościowy", section="badanie")
items = [
    "1. wywiad: objawy, zaostrzenia, leczenie, dotychczasowe ACT",
    "2. badanie przedmiotowe: oddech, kaszel, wydzielina, postawa",
    "3. ocena funkcji płuc: spirometria, saturacja, ewentualnie obrazowanie",
    "4. ocena wydolności: CPET lub testy terenowe",
    "5. ocena jakości życia, adherencji i barier",
    "6. synteza ryzyka i priorytetów terapii",
]
bullet_box(s, items, x=1.0, y=1.55, w=11.1, h=5.0, font_size=18)
add_footer(s, 8, "Hebestreit et al., 2015")

# 9
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Wywiad fizjoterapeutyczny", section="badanie")
two_col(
    s,
    "Pytania kluczowe",
    [
        "kaszel: częstość, produktywność, pora dnia",
        "ilość, barwa i lepkość wydzieliny",
        "duszność, tolerancja schodów, szkoły/pracy/sportu",
        "zaostrzenia: liczba, hospitalizacje, antybiotyki",
        "aktualne leki wziewne i kolejność terapii",
    ],
    "Sygnały alarmowe",
    [
        "nagły spadek saturacji lub narastająca duszność",
        "gorączka, ból opłucnowy, znaczne osłabienie",
        "krwioplucie, podejrzenie odmy",
        "spadek masy ciała lub odwodnienie",
        "brak tolerancji dotychczasowej techniki ACT",
    ],
)
add_footer(s, 9)

# 10
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Badanie przedmiotowe", section="badanie")
bullet_box(
    s,
    [
        "obserwacja toru oddechowego, częstości oddechów i pracy dodatkowych mięśni oddechowych",
        "ocena kaszlu: efektywność, kontrola, technika huff/FET",
        "osłuchiwanie: świsty, furczenia, trzeszczenia, asymetria szmeru oddechowego",
        "ocena klatki piersiowej: ruchomość, kifoza piersiowa, ustawienie barków",
        "pomiar SpO2 i tętna w spoczynku oraz po wysiłku",
        "ocena zmęczenia, bólu, tolerancji pozycji drenażowych",
    ],
    x=0.9,
    y=1.55,
    w=11.4,
    h=5.15,
    font_size=17,
)
add_footer(s, 10)

# 11
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Badania czynnościowe układu oddechowego", section="badanie")
two_col(
    s,
    "Spirometria i parametry",
    [
        "FEV1 % należnej: podstawowy marker obturacji i progresji",
        "FVC, FEV1/FVC, przepływy wydechowe",
        "trend jest ważniejszy niż pojedynczy pomiar",
        "pomiar przed/po interwencji może pokazać odpowiedź na ACT",
    ],
    "Uzupełnienie oceny",
    [
        "SpO2 w spoczynku, podczas ACT i wysiłku",
        "badanie mikrobiologiczne plwociny",
        "obrazowanie płuc według decyzji zespołu",
        "kontrola objawów zaostrzenia i działań niepożądanych",
    ],
)
add_footer(s, 11, "Castellani et al., 2018")

# 12
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Ocena wydzieliny i efektywności kaszlu", section="badanie")
bullet_box(
    s,
    [
        "ilość: subiektywna skala, objętość lub masa mokrej plwociny po sesji",
        "jakość: lepkość, barwa, obecność krwi, zmiana zapachu",
        "lokalizacja zalegania: osłuchowo, objawowo, według preferowanych pozycji",
        "efektywność: czy huff/FET przesuwa wydzielinę bez nadmiernego skurczu oskrzeli",
        "reakcja po sesji: duszność, zmęczenie, SpO2, uczucie oczyszczenia",
    ],
    x=1.0,
    y=1.65,
    w=10.9,
    h=4.8,
    font_size=18,
)
add_footer(s, 12, "McIlwaine et al., 2019")

# 13
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Ocena wydolności i aktywności fizycznej", section="badanie")
two_col(
    s,
    "Złoty standard",
    [
        "CPET: szczytowy pobór tlenu i maksymalna moc pracy",
        "monitorowanie SpO2, tętna, objawów i wentylacji",
        "u osób >=10 lat preferowany protokół cykloergometru Godfreya",
        "identyfikacja ograniczeń i bezpieczeństwa wysiłku",
    ],
    "Gdy CPET niedostępny",
    [
        "test marszowy 6-minutowy lub shuttle test",
        "skale duszności i zmęczenia, np. Borg",
        "dzienniczek aktywności lub krokomierz",
        "ocena preferencji: sport, zabawa, trening domowy",
    ],
)
add_footer(s, 13, "Hebestreit et al., 2015")

# 14
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Ocena jakości życia i obciążenia leczeniem", section="badanie")
bullet_box(
    s,
    [
        "CFQ-R: kwestionariusz swoisty dla mukowiscydozy, obejmuje m.in. objawy oddechowe i obciążenie leczeniem",
        "wywiad o adherencji: czas terapii, zapominanie, wstyd, koszty, trudność techniki",
        "ocena funkcjonowania: sen, szkoła/praca, aktywność społeczna, lęk przed zaostrzeniem",
        "wspólne decyzje zwiększają szansę, że pacjent realnie wykona plan",
    ],
    x=0.95,
    y=1.65,
    w=11.1,
    h=4.6,
    font_size=18,
)
add_footer(s, 14, "Gomez et al., 2018; Lannefors et al., 2022")

# 15
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Diagnoza fizjoterapeutyczna: przykładowa synteza", section="badanie")
callout(s, "problem główny", 0.8, 1.5, 2.2, 0.6, PALE, TEAL)
bullet_box(
    s,
    [
        "zaleganie gęstej wydzieliny z produktywnym kaszlem rano i po wysiłku",
        "obniżona tolerancja wysiłku, szybka męczliwość i okresowe spadki SpO2",
        "ograniczona ruchomość klatki piersiowej oraz tendencja do kifotyzacji",
        "duże obciążenie czasowe leczenia i nieregularne wykonywanie ACT",
    ],
    x=1.05,
    y=2.25,
    w=10.8,
    h=3.0,
    font_size=18,
)
callout(s, "wniosek: potrzebny program łączący oczyszczanie, trening, edukację i monitoring", 1.25, 5.55, 10.8, 0.7, RGBColor(255, 244, 235), ORANGE, 15)
add_footer(s, 15)

# 16
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Przykładowy pacjent do planu terapii", section="plan terapii")
two_col(
    s,
    "Założenia kliniczne",
    [
        "pacjent nastoletni lub młody dorosły w stabilnym okresie choroby",
        "rozpoznana mukowiscydoza, kaszel produktywny",
        "FEV1 łagodnie/umiarkowanie obniżone, bez ostrego krwioplucia",
        "samodzielny, ale z problemem regularności terapii",
    ],
    "Priorytety",
    [
        "codzienne oczyszczanie oskrzeli możliwe do wykonania w domu",
        "zwiększenie tolerancji wysiłku i aktywności rekreacyjnej",
        "nauka kontroli kaszlu i higieny sprzętu",
        "wczesne rozpoznanie zaostrzenia",
    ],
)
add_footer(s, 16)

# 17
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Algorytm planowania sesji fizjoterapii", section="plan terapii")
for i, (text, col) in enumerate(
    [
        ("1. inhalacja\njeśli zalecona", TEAL),
        ("2. ACT\nPEP / OPEP / ACBT", CYAN),
        ("3. huff + kaszel\nkontrolowany", GREEN),
        ("4. ćwiczenia\noddechowe i mobilność", ORANGE),
        ("5. trening\nwytrzymałościowy / siłowy", RGBColor(117, 100, 170)),
        ("6. zapis efektów\nobjawy, plwocina, SpO2", RGBColor(47, 121, 180)),
    ]
):
    x = 0.65 + (i % 3) * 4.15
    y = 1.65 + (i // 3) * 2.15
    callout(s, text, x, y, 3.35, 1.25, RGBColor(245, 250, 251), col, 14)
bullet_box(
    s,
    ["Kolejność należy dopasować do zaleceń lekarskich, tolerancji wysiłku i pory dnia z największym zaleganiem wydzieliny."],
    x=1.0,
    y=6.0,
    w=11.2,
    h=0.45,
    font_size=15,
)
add_footer(s, 17)

# 18
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Techniki oczyszczania dróg oddechowych (ACT)", section="interwencje")
two_col(
    s,
    "Techniki bez urządzeń",
    [
        "ACBT: kontrola oddechu, ćwiczenia rozszerzające, FET/huff",
        "drenaż autogeniczny: kontrola objętości płuc i przepływu",
        "pozycje ułatwiające wentylację i odpływ wydzieliny",
        "kaszel kontrolowany z oszczędzaniem energii",
    ],
    "Techniki z urządzeniami",
    [
        "PEP maska/ustnik: dodatnie ciśnienie wydechowe",
        "oscylacyjny PEP: Flutter, Acapella, Aerobika",
        "HFCWO: kamizelka oscylacyjna u wybranych pacjentów",
        "dobór zależy od efektu, tolerancji i preferencji",
    ],
)
add_footer(s, 18, "Flume et al., 2009")

# 19
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "PEP i oscylacyjny PEP: praktyczne zasady", section="interwencje")
bullet_box(
    s,
    [
        "cel: utrzymać drożność małych dróg oddechowych i przesuwać powietrze za czop śluzowy",
        "typowy schemat: serie spokojnych wydechów przez opór, przerwy na huff i kaszel",
        "monitorować: duszność, zawroty głowy, skurcz oskrzeli, SpO2, zmęczenie",
        "PEP bywa preferowany przez pacjentów i w długich badaniach ograniczał zaostrzenia względem HFCWO",
        "sprzęt wymaga instruktażu, kontroli oporu oraz regularnej dezynfekcji",
    ],
    x=0.95,
    y=1.55,
    w=11.0,
    h=5.1,
    font_size=17,
)
add_footer(s, 19, "McIlwaine et al., 2013; McIlwaine et al., 2019")

# 20
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "ACBT i drenaż autogeniczny", section="interwencje")
two_col(
    s,
    "ACBT",
    [
        "kontrola oddechu zmniejsza skurcz oskrzeli i zmęczenie",
        "głębsze wdechy z pauzą poprawiają wentylację obwodową",
        "FET/huff przesuwa wydzielinę z mniejszym zapadaniem oskrzeli niż gwałtowny kaszel",
        "łatwy do łączenia z pozycjami i inhalacją",
    ],
    "Drenaż autogeniczny",
    [
        "oddychanie na niskich, średnich i wysokich objętościach płuc",
        "wymaga większej świadomości oddechu i treningu techniki",
        "przydatny u pacjentów, którzy chcą techniki dyskretnej i bez sprzętu",
        "efekt oceniamy po wydzielinie, objawach i tolerancji",
    ],
)
add_footer(s, 20)

# 21
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Inhalacje, kolejność i higiena", section="interwencje")
bullet_box(
    s,
    [
        "fizjoterapeuta powinien znać zalecony przez lekarza schemat leków wziewnych i mukolitycznych",
        "często ACT planuje się po inhalacji ułatwiającej upłynnienie wydzieliny; bronchodilatator przed wysiłkiem, jeśli zalecony",
        "nie mieszać dowolnie kolejności leków bez zaleceń zespołu prowadzącego",
        "nebulizatory, maski PEP i ustniki: mycie, suszenie i dezynfekcja według instrukcji",
        "kontrola zakażeń: własny sprzęt, wentylowane miejsce, unikanie kontaktu z innymi pacjentami CF",
    ],
    x=0.9,
    y=1.55,
    w=11.3,
    h=5.1,
    font_size=17,
)
add_footer(s, 21, "Castellani et al., 2018")

# 22
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Trening fizyczny jako element terapii", section="interwencje")
two_col(
    s,
    "Cele treningu",
    [
        "poprawa wydolności krążeniowo-oddechowej",
        "wspomaganie oczyszczania wydzieliny przez wentylację i huff",
        "wzrost siły mięśniowej i tolerancji codziennych aktywności",
        "korzyści psychologiczne i społeczne",
    ],
    "Dobór obciążeń",
    [
        "aerobowy: marsz, rower, pływanie, gry, interwały",
        "siłowy: duże grupy mięśniowe, technika i oddech bez parcia",
        "intensywność wg CPET/testów terenowych, SpO2 i skali Borga",
        "nawodnienie i podaż energii muszą odpowiadać większemu wysiłkowi",
    ],
)
add_footer(s, 22, "Williams et al., 2022; Radtke et al., 2022")

# 23
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Ćwiczenia oddechowe, posturalne i mobilizacyjne", section="interwencje")
bullet_box(
    s,
    [
        "mobilizacja odcinka piersiowego: wyprost, rotacje, otwieranie klatki piersiowej",
        "rozciąganie mięśni piersiowych, zginaczy bioder i obręczy barkowej",
        "ćwiczenia oddechowe z akcentem na dolnożebrowy tor oddechowy i kontrolę wydechu",
        "trening mięśni posturalnych oraz stabilizacji łopatki",
        "łączenie mobilizacji z fazą wdechu i huffem może zwiększać komfort oczyszczania",
    ],
    x=0.95,
    y=1.6,
    w=11.0,
    h=4.9,
    font_size=18,
)
add_footer(s, 23, "Malkoc et al., 2021")

# 24
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Modyfikacja terapii w zaostrzeniu", section="bezpieczeństwo")
two_col(
    s,
    "Co zwykle zwiększamy?",
    [
        "częstość ACT, jeśli pacjent ma więcej wydzieliny",
        "monitorowanie SpO2, tętna, duszności i zmęczenia",
        "pozycje ułatwiające wentylację i odpoczynek",
        "krótsze, częstsze sesje zamiast jednej długiej",
    ],
    "Kiedy przerwać i skonsultować?",
    [
        "krwioplucie, podejrzenie odmy, ból w klatce piersiowej",
        "istotna desaturacja lub narastająca duszność",
        "gorączka z pogorszeniem stanu ogólnego",
        "brak tolerancji dotychczasowego oporu/pozycji",
    ],
)
add_footer(s, 24)

# 25
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Edukacja pacjenta i rodziny", section="edukacja")
bullet_box(
    s,
    [
        "pacjent rozumie: po co wykonuje ACT, kiedy ją zwiększyć i kiedy zgłosić się do lekarza",
        "nauka huff, kaszlu kontrolowanego, ustawienia urządzenia i higieny sprzętu",
        "plan dnia: terapia wpisana w realistyczne pory, np. rano i wieczorem",
        "strategie adherencji: krótkie cele, dzienniczek, aplikacja, wsparcie rodziny",
        "aktywność fizyczna jako styl życia, nie tylko zadanie rehabilitacyjne",
    ],
    x=0.95,
    y=1.55,
    w=11.1,
    h=5.05,
    font_size=18,
)
add_footer(s, 25, "Williams et al., 2022")

# 26
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Monitorowanie efektów terapii", section="monitorowanie")
two_col(
    s,
    "Wskaźniki krótkoterminowe",
    [
        "łatwiejsze odkrztuszanie i mniejsza duszność po sesji",
        "stabilna SpO2 i tętno podczas ACT",
        "mniej uczucia zalegania, poprawa osłuchowa",
        "pacjent potrafi sam wykonać technikę",
    ],
    "Wskaźniki długoterminowe",
    [
        "trend FEV1 i liczba zaostrzeń",
        "wydolność: CPET / 6MWT / shuttle test",
        "CFQ-R: objawy oddechowe i obciążenie leczeniem",
        "regularność terapii i aktywności fizycznej",
    ],
)
add_footer(s, 26, "Hebestreit et al., 2015; Gomez et al., 2018")

# 27
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Przykładowy tygodniowy plan terapii domowej", section="plan terapii")
rows, cols = 6, 4
table = s.shapes.add_table(rows, cols, Inches(0.55), Inches(1.45), Inches(12.25), Inches(4.65)).table
widths = [2.1, 3.35, 3.35, 3.45]
for i, width in enumerate(widths):
    table.columns[i].width = Inches(width)
headers = ["Element", "Pon.-pt.", "Weekend", "Kontrola"]
data = [
    ["ACT", "PEP/OPEP 2 x dziennie, 15-25 min", "1-2 x dziennie wg objawów", "plwocina, duszność, SpO2"],
    ["Huff/kaszel", "po każdej serii ACT i po wysiłku", "jak w dni robocze", "efektywność bez wyczerpania"],
    ["Trening aerobowy", "3-5 sesji: marsz/rower/interwały", "aktywność rekreacyjna", "Borg, tętno, SpO2"],
    ["Siła i postawa", "2-3 sesje: tułów, obręcz barkowa, nogi", "mobilizacja i rozciąganie", "ból, technika oddechu"],
    ["Edukacja", "dzienniczek i higiena sprzętu", "przegląd planu tygodnia", "adherencja, bariery"],
]
for c, htxt in enumerate(headers):
    cell = table.cell(0, c)
    cell.text = htxt
    cell.fill.solid()
    cell.fill.fore_color.rgb = TEAL
    for p in cell.text_frame.paragraphs:
        p.font.color.rgb = WHITE
        p.font.bold = True
        p.font.size = Pt(11)
for r_i, row in enumerate(data, start=1):
    for c, val in enumerate(row):
        cell = table.cell(r_i, c)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = LIGHT if r_i % 2 else WHITE
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(10)
            p.font.color.rgb = DARK
callout(s, "Plan wymaga akceptacji zespołu prowadzącego i aktualizacji przy zaostrzeniu.", 1.15, 6.25, 11.0, 0.55, RGBColor(255, 244, 235), ORANGE, 13)
add_footer(s, 27)

# 28
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Bezpieczeństwo i przeciwwskazania względne", section="bezpieczeństwo")
bullet_box(
    s,
    [
        "dostosować intensywność przy gorączce, odwodnieniu, znacznym zmęczeniu lub bólu",
        "unikać technik nasilających krwioplucie; przy istotnym krwiopluciu wymagana pilna konsultacja",
        "podejrzenie odmy: nie wykonywać intensywnego ACT ani wysiłku do oceny lekarskiej",
        "desaturacja, ból w klatce piersiowej lub omdlenie podczas treningu wymagają przerwania ćwiczeń",
        "ryzyko zakażeń krzyżowych: pacjenci z CF nie powinni współdzielić sprzętu ani ćwiczyć blisko siebie",
    ],
    x=0.95,
    y=1.55,
    w=11.1,
    h=5.0,
    font_size=18,
)
add_footer(s, 28)

# 29
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Najważniejsze wnioski", section="podsumowanie")
bullet_box(
    s,
    [
        "fizjoterapia w mukowiscydozie jest leczeniem długoterminowym, regularnym i indywidualizowanym",
        "badanie pacjenta obejmuje objawy, wydzielinę, funkcję płuc, wydolność, postawę, jakość życia i adherencję",
        "nie ma jednej najlepszej techniki ACT dla wszystkich; skuteczny jest plan, który pacjent wykonuje prawidłowo i systematycznie",
        "aktywność fizyczna i trening są ważnym uzupełnieniem oczyszczania dróg oddechowych",
        "monitorowanie trendów FEV1, zaostrzeń, tolerancji wysiłku i CFQ-R pozwala modyfikować terapię",
    ],
    x=0.9,
    y=1.55,
    w=11.25,
    h=5.2,
    font_size=18,
)
add_footer(s, 29)

# 30
s = prs.slides.add_slide(prs.slide_layouts[6])
add_title(s, "Bibliografia", section="literatura")
refs = [
    "Flume PA et al. Cystic fibrosis pulmonary guidelines: airway clearance therapies. Respir Care. 2009;54(4):522-537. PMID: 19327189.",
    "McIlwaine MP et al. High frequency chest wall oscillation versus PEP mask in cystic fibrosis. Thorax. 2013;68(8):746-751. PMID: 23407019.",
    "McIlwaine M, Button B, Nevitt SJ. Positive expiratory pressure physiotherapy for airway clearance in people with cystic fibrosis. Cochrane Database Syst Rev. 2019;11:CD003147. PMID: 31774149.",
    "Hebestreit H et al. Statement on Exercise Testing in Cystic Fibrosis. Respiration. 2015;90(4):332-351. PMID: 26352941.",
    "Castellani C et al. ECFS best practice guidelines: the 2018 revision. J Cyst Fibros. 2018;17(2):153-178.",
    "Williams CA et al. The Exeter Activity Unlimited statement on physical activity and exercise for cystic fibrosis. Chronic Respir Dis. 2022;19:14799731221121670.",
    "Radtke T et al. Physical activity and exercise training in cystic fibrosis. Cochrane Database Syst Rev. 2022;8:CD002768.",
    "Gomez CC et al. Development and electronic validation of the revised Cystic Fibrosis Questionnaire. J Cyst Fibros. 2018;17(1):97-102. PMID: 29157922.",
    "Malkoc M et al. Combining postural exercises with chest physiotherapy in cystic fibrosis: randomized-controlled trial. Turk J Phys Med Rehabil. 2021;67(1):91-100. PMID: 33948549.",
]
bullet_box(s, refs, x=0.65, y=1.25, w=12.1, h=5.65, font_size=10)
add_footer(s, 30, "PubMed / Google Scholar")

prs.save(OUTFILE)
print(OUTFILE)
