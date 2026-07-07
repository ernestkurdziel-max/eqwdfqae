#!/usr/bin/env python3
"""Generuje prezentację PPTX na podstawie pracy o fizjoterapii w mukowiscydozie."""

from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PRIMARY = RGBColor(0x0B, 0x4F, 0x6C)      # ciemny turkus
ACCENT = RGBColor(0x14, 0x8F, 0xB5)       # jasny turkus
DARK = RGBColor(0x22, 0x2B, 0x33)
LIGHT = RGBColor(0xF2, 0xF7, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

blank = prs.slide_layouts[6]


def add_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_bar(slide, color=PRIMARY, height=Inches(1.25)):
    bar = slide.shapes.add_shape(1, 0, 0, SW, height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()
    bar.shadow.inherit = False
    return bar


def add_text(slide, left, top, width, height, text, size, color=DARK,
             bold=False, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return tb


def title_slide(title, subtitle, footer):
    slide = prs.slides.add_slide(blank)
    add_bg(slide, PRIMARY)
    accent = slide.shapes.add_shape(1, 0, Inches(4.6), SW, Inches(0.12))
    accent.fill.solid(); accent.fill.fore_color.rgb = ACCENT; accent.line.fill.background()
    accent.shadow.inherit = False
    add_text(slide, Inches(0.9), Inches(2.0), Inches(11.5), Inches(2.4), title,
             40, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, Inches(1.2), Inches(4.8), Inches(10.9), Inches(1.2), subtitle,
             24, RGBColor(0xCF, 0xEA, 0xF2), align=PP_ALIGN.CENTER)
    add_text(slide, Inches(1.2), Inches(6.6), Inches(10.9), Inches(0.6), footer,
             14, RGBColor(0xB8, 0xD8, 0xE2), align=PP_ALIGN.CENTER)
    return slide


def content_slide(title, bullets, kicker=None):
    slide = prs.slides.add_slide(blank)
    add_bg(slide, LIGHT)
    add_bar(slide)
    if kicker:
        add_text(slide, Inches(0.7), Inches(0.18), Inches(12), Inches(0.35),
                 kicker.upper(), 12, RGBColor(0xBF, 0xE3, 0xEE), bold=True)
    add_text(slide, Inches(0.7), Inches(0.42), Inches(12), Inches(0.8), title,
             28, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)

    body = slide.shapes.add_textbox(Inches(0.8), Inches(1.55), Inches(11.8), Inches(5.6))
    tf = body.text_frame
    tf.word_wrap = True
    first = True
    for item in bullets:
        level = 0
        text = item
        if isinstance(item, tuple):
            level, text = item
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(8)
        run = p.add_run()
        if level == 0:
            run.text = "▪  " + text
            run.font.size = Pt(19)
            run.font.bold = True
            run.font.color.rgb = PRIMARY
        else:
            run.text = "–  " + text
            run.font.size = Pt(16)
            run.font.color.rgb = DARK
    return slide


def section_slide(number, title):
    slide = prs.slides.add_slide(blank)
    add_bg(slide, PRIMARY)
    add_text(slide, Inches(0.9), Inches(2.4), Inches(2.5), Inches(2), number,
             90, ACCENT, bold=True)
    add_text(slide, Inches(3.2), Inches(2.7), Inches(9.2), Inches(2), title,
             34, WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    return slide


# --- Slajd tytułowy ---
title_slide(
    "Zasady planowania i programowania fizjoterapii\npacjentów z chorobami układu oddechowego",
    "Mukowiscydoza (zwłóknienie torbielowate)",
    "Opracowanie o charakterze przeglądowo-edukacyjnym",
)

# --- Plan prezentacji ---
content_slide("Plan prezentacji", [
    "Charakterystyka mukowiscydozy",
    "Rola i cele fizjoterapii",
    "Ocena (diagnostyka) fizjoterapeutyczna",
    "Zasady planowania fizjoterapii",
    "Programowanie – metody i techniki",
    "Dostosowanie do wieku i stanu klinicznego",
    "Monitorowanie, bezpieczeństwo i kontrola zakażeń",
    "Podsumowanie i wnioski",
], kicker="Agenda")

# --- Sekcja 1 ---
section_slide("1", "Charakterystyka mukowiscydozy")

content_slide("Czym jest mukowiscydoza?", [
    "Najczęstsza letalna choroba genetyczna rasy białej",
    "Dziedziczenie autosomalne recesywne",
    "Przyczyna: mutacja genu CFTR (chromosom 7); najczęstsza – F508del",
    (1, "CFTR koduje kanał chlorkowy w błonach komórek nabłonka"),
    "Skutek: zaburzony transport jonów Cl⁻ i Na⁺ → gęsta, lepka wydzielina",
    "Choroba wielonarządowa – dominują objawy oddechowe",
], kicker="Definicja i etiologia")

content_slide("Patofizjologia – „błędne koło”", [
    "Odwodnienie warstwy płynu okołorzęskowego",
    "Upośledzenie klirensu śluzowo-rzęskowego",
    "Zaleganie wydzieliny → kolonizacja bakteryjna",
    (1, "m.in. Pseudomonas aeruginosa, Staphylococcus aureus"),
    "Przewlekłe zakażenie → stan zapalny → uszkodzenie oskrzeli",
    "Efekt: rozstrzenie oskrzeli, obturacja, niewydolność oddechowa",
], kicker="Mechanizm")

content_slide("Objawy i diagnostyka", [
    "Objawy oddechowe: przewlekły kaszel, ropna wydzielina, duszność",
    (1, "nawracające zakażenia, obturacja, spadek tolerancji wysiłku"),
    "Objawy pozapłucne: niewydolność trzustki, niedożywienie, CFRD, osteoporoza",
    "Diagnostyka:",
    (1, "przesiew noworodkowy (IRT)"),
    (1, "test potowy (chlorki ≥ 60 mmol/l)"),
    (1, "badania genetyczne (mutacje CFTR)"),
], kicker="Obraz kliniczny")

# --- Sekcja 2 ---
section_slide("2", "Rola i cele fizjoterapii")

content_slide("Planowanie a programowanie", [
    "PLANOWANIE – poziom strategiczny",
    (1, "cele, kierunki i ramy terapii (po co? dokąd? w jakim czasie?)"),
    "PROGRAMOWANIE – poziom operacyjny",
    (1, "dobór metod, technik i parametrów obciążenia (co? jak? jak często?)"),
    "Oba etapy sprzężone zwrotnie – wyniki modyfikują plan",
], kicker="Pojęcia kluczowe")

content_slide("Cele fizjoterapii w mukowiscydozie", [
    "Utrzymanie drożności dróg oddechowych (usuwanie wydzieliny)",
    "Poprawa wentylacji i wymiany gazowej",
    "Zapobieganie zakażeniom i powikłaniom",
    "Utrzymanie/poprawa wydolności fizycznej i tolerancji wysiłku",
    "Prawidłowa postawa i ruchomość klatki piersiowej",
    "Edukacja i samodzielność terapeutyczna pacjenta",
], kicker="Cele")

content_slide("Ogólne zasady postępowania", [
    "Indywidualizacja",
    "Systematyczność i ciągłość (terapia dożywotnia)",
    "Stopniowanie obciążeń (progresja)",
    "Wczesne rozpoczęcie – od momentu rozpoznania",
    "Kompleksowość (zespół wielospecjalistyczny)",
    "Aktywny udział pacjenta",
    "Ocena efektów i bezpieczeństwo",
], kicker="Zasady")

# --- Sekcja 3 ---
section_slide("3", "Ocena fizjoterapeutyczna")

content_slide("Diagnostyka funkcjonalna", [
    "Wywiad: przebieg choroby, wydzielina, duszność, aktywność, odżywienie",
    "Badanie przedmiotowe: oglądanie, palpacja, osłuchiwanie, postawa",
    "Badania czynnościowe płuc:",
    (1, "spirometria (FEV1, FVC, FEV1/FVC), pletyzmografia, pulsoksymetria"),
    "Ocena wydolności: test 6-minutowego marszu (6MWT), CPET (VO₂peak)",
    "Skale: duszności (Borg, mMRC), jakości życia (CFQ-R)",
], kicker="Ocena wyjściowa")

# --- Sekcja 4 ---
section_slide("4", "Zasady planowania")

content_slide("Etapy planowania i cele SMART", [
    "Etapy: analiza danych → diagnoza → cele → strategia → ramy → kryteria oceny",
    "Cele wg reguły SMART:",
    (1, "Konkretne, Mierzalne, Osiągalne, Istotne, Określone w czasie"),
    "Przykłady celów:",
    (1, "krótkoterminowy: opanowanie techniki ACBT (2 tyg.)"),
    (1, "średnioterminowy: wzrost dystansu 6MWT (8–12 tyg.)"),
    (1, "długoterminowy: stabilizacja FEV1 i aktywności fizycznej"),
], kicker="Strategia")

content_slide("Indywidualizacja i priorytety", [
    "Dostosowanie do wieku i fazy choroby",
    "Uwzględnienie chorób współistniejących (CFRD, osteoporoza)",
    "Preferencje i styl życia → przestrzeganie zaleceń (compliance)",
    "Hierarchizacja priorytetów:",
    (1, "zaostrzenie → intensywne oczyszczanie oskrzeli"),
    (1, "okres stabilny → wydolność, trening, profilaktyka"),
], kicker="Dopasowanie")

# --- Sekcja 5 ---
section_slide("5", "Programowanie – metody i techniki")

content_slide("Techniki oczyszczania oskrzeli (ACT)", [
    "Cykl aktywnego oddychania (ACBT) – kontrola, rozprężanie, huffing",
    "Drenaż autogeniczny (AD)",
    "Urządzenia PEP (dodatnie ciśnienie wydechowe)",
    "Oscylacyjne PEP – Flutter, Acapella, RC-Cornet",
    "Drenaż ułożeniowy, oklepywanie, wibracje (ostrożnie)",
    "HFCWO – kamizelka oscylacyjna",
    (1, "1–3× dziennie, sesje ok. 15–30 min; brak jednej „najlepszej” techniki"),
], kicker="Filar terapii")

content_slide("Trening fizyczny – schemat FITT", [
    "Wyższa wydolność tlenowa → lepsze rokowanie",
    "Rodzaje: wytrzymałościowy (aerobowy), oporowy, gibkościowy",
    "F – częstotliwość: aktywność w większości dni, trening 3–5×/tydz.",
    "I – intensywność: indywidualnie (tętno, SpO₂, skala Borga)",
    "T – czas: ok. 20–40 min",
    "T – rodzaj: połączenie treningu aerobowego i oporowego",
    (1, "monitoring SpO₂/tętna, nawodnienie i suplementacja soli"),
], kicker="Wydolność")

content_slide("Terapia wziewna i edukacja", [
    "Inhalacje wspierają oczyszczanie oskrzeli",
    "Właściwa sekwencja:",
    (1, "lek rozszerzający oskrzela → mukolityk/sól hipertoniczna →"),
    (1, "techniki oczyszczania → antybiotyk wziewny (jeśli zlecony)"),
    "Edukacja pacjenta i rodziny: technika, obsługa sprzętu, higiena",
    "Budowanie motywacji i nawyku codziennej terapii",
], kicker="Wsparcie")

# --- Sekcja 6 ---
section_slide("6", "Wiek i stan kliniczny")

content_slide("Dostosowanie programu", [
    "Niemowlęta/małe dzieci: techniki bierne, zabawa, edukacja rodziców",
    "Dzieci starsze/młodzież: ACBT, AD, PEP/OPEP, sport, samodzielność",
    "Dorośli: samodzielna terapia, wydolność, zarządzanie chorobą",
    "Okres stabilny: utrzymanie funkcji płuc i wydolności",
    "Zaostrzenie: intensyfikacja oczyszczania, łagodny trening",
    "Choroba zaawansowana: tlenoterapia, NIV, prehabilitacja, opieka paliatywna",
], kicker="Personalizacja")

# --- Sekcja 7 ---
section_slide("7", "Monitorowanie i bezpieczeństwo")

content_slide("Monitorowanie efektów", [
    "Spirometria (FEV1, FVC) – progresja choroby",
    "Testy wydolnościowe (6MWT, CPET)",
    "Ocena objawów: kaszel, wydzielina, duszność, SpO₂",
    "Kwestionariusze jakości życia (CFQ-R)",
    "Ocena przestrzegania zaleceń i stanu odżywienia",
    (1, "wyniki → modyfikacja technik, obciążeń i celów"),
], kicker="Weryfikacja")

content_slide("Bezpieczeństwo i kontrola zakażeń", [
    "Ostrożność / przeciwwskazania:",
    (1, "krwioplucie, odma opłucnowa, refluks, hipoksemia, osteoporoza"),
    "Kontrola zakażeń krzyżowych (P. aeruginosa, B. cepacia):",
    (1, "segregacja pacjentów, higiena rąk"),
    (1, "dezynfekcja i indywidualizacja sprzętu"),
    (1, "terapia indywidualna, telerehabilitacja"),
], kicker="Ryzyko")

# --- Podsumowanie ---
section_slide("8", "Podsumowanie i wnioski")

content_slide("Wnioski", [
    "Fizjoterapia to nieodłączny, dożywotni element leczenia mukowiscydozy",
    "Proces: ocena → planowanie → programowanie → realizacja → monitorowanie",
    "Fundament: techniki oczyszczania oskrzeli + regularny trening fizyczny",
    "Kluczowe: indywidualizacja, systematyczność, wczesne rozpoczęcie",
    "Program elastycznie dostosowany do wieku i fazy choroby",
    "Efekt: spowolnienie progresji choroby płuc i poprawa jakości życia",
], kicker="Kluczowe przesłania")

# --- Slajd końcowy ---
end = prs.slides.add_slide(blank)
add_bg(end, PRIMARY)
add_text(end, Inches(1), Inches(2.9), Inches(11.3), Inches(1.5), "Dziękuję za uwagę",
         44, WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
add_text(end, Inches(1), Inches(4.5), Inches(11.3), Inches(0.8),
         "Pytania i dyskusja", 22, RGBColor(0xCF, 0xEA, 0xF2), align=PP_ALIGN.CENTER)

prs.save("prezentacja_fizjoterapia_mukowiscydoza.pptx")
print("Zapisano prezentację.")
