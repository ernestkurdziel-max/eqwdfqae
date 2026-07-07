# -*- coding: utf-8 -*-
"""
Generator prezentacji PowerPoint:
"Zasady planowania i programowania fizjoterapii pacjentow
 z chorobami ukladu oddechowego - Astma oskrzelowa"

Przedmiot: Kliniczne podstawy fizjoterapii w pulmonologii
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------------------
# Paleta kolorow (motyw pulmonologiczny - odcienie blekitu i zieleni)
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x0B, 0x3C, 0x5D)      # ciemny granat - naglowki
TEAL = RGBColor(0x14, 0x8F, 0x9E)      # morski - akcenty
LIGHT_TEAL = RGBColor(0xE3, 0xF2, 0xF4)  # jasny blekit - tla ramek
ACCENT = RGBColor(0x2A, 0x9D, 0x8F)    # zielono-morski
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x22, 0x2B, 0x33)
GREY_TEXT = RGBColor(0x4A, 0x55, 0x60)
LIGHT_BG = RGBColor(0xF4, 0xF9, 0xFA)
CORAL = RGBColor(0xE7, 0x6F, 0x51)     # akcent ostrzegawczy

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H

BLANK = prs.slide_layouts[6]


# ---------------------------------------------------------------------------
# Funkcje pomocnicze
# ---------------------------------------------------------------------------
def add_slide():
    return prs.slides.add_slide(BLANK)


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def add_round_rect(slide, x, y, w, h, color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(1.25)
    shape.shadow.inherit = False
    return shape


def add_text(slide, x, y, w, h, text, size=18, color=DARK_TEXT, bold=False,
             italic=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             font="Calibri", line_spacing=1.0):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = font
    return tb


def add_bullets(slide, x, y, w, h, items, size=18, color=DARK_TEXT,
                bullet_color=TEAL, line_spacing=1.12, space_after=8,
                font="Calibri"):
    """items: list of (text, level) tuples or plain strings (level 0)."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    first = True
    for item in items:
        if isinstance(item, tuple):
            txt, level = item
        else:
            txt, level = item, 0
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.line_spacing = line_spacing
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        # marker
        marker = "•  " if level == 0 else "–  "
        r1 = p.add_run()
        r1.text = marker
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = bullet_color if level == 0 else ACCENT
        r1.font.name = font
        r2 = p.add_run()
        r2.text = txt
        r2.font.size = Pt(size if level == 0 else size - 1)
        r2.font.color.rgb = color
        r2.font.name = font
    return tb


slide_no = [0]


def add_header(slide, title, subtitle=None):
    """Standardowy naglowek slajdu tresciowego."""
    add_rect(slide, 0, 0, SLIDE_W, Inches(1.15), NAVY)
    add_rect(slide, 0, Inches(1.15), SLIDE_W, Inches(0.08), TEAL)
    # pasek akcentu
    add_rect(slide, 0, 0, Inches(0.22), Inches(1.15), TEAL)
    add_text(slide, Inches(0.55), Inches(0.14), Inches(11.8), Inches(0.9),
             title, size=27, color=WHITE, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
    if subtitle:
        add_text(slide, Inches(0.57), Inches(0.78), Inches(11.8), Inches(0.35),
                 subtitle, size=13, color=LIGHT_TEAL, italic=True)


def add_footer(slide):
    slide_no[0] += 1
    add_text(slide, Inches(0.4), Inches(7.02), Inches(9.0), Inches(0.4),
             "Astma oskrzelowa – planowanie i programowanie fizjoterapii",
             size=9, color=GREY_TEXT)
    add_text(slide, Inches(12.2), Inches(7.02), Inches(0.9), Inches(0.4),
             str(slide_no[0]), size=10, color=GREY_TEXT, align=PP_ALIGN.RIGHT)


def content_slide(title, subtitle=None, bg=LIGHT_BG):
    s = add_slide()
    set_bg(s, bg)
    add_header(s, title, subtitle)
    add_footer(s)
    return s


def card(slide, x, y, w, h, title, body_items, title_color=WHITE,
         head_color=TEAL, body_bg=WHITE, body_size=15):
    """Karta z naglowkiem i lista punktow."""
    add_round_rect(slide, x, y, w, h, body_bg, line_color=RGBColor(0xD5, 0xE3, 0xE6))
    head = add_round_rect(slide, x, y, w, Inches(0.55), head_color)
    add_text(slide, x + Inches(0.15), y, w - Inches(0.3), Inches(0.55),
             title, size=15, color=title_color, bold=True,
             anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(slide, x + Inches(0.22), y + Inches(0.7),
                w - Inches(0.44), h - Inches(0.85), body_items,
                size=body_size, space_after=5, line_spacing=1.05)


# ===========================================================================
# SLAJD 1 - TYTULOWY
# ===========================================================================
s = add_slide()
set_bg(s, NAVY)
# ozdobne pasy
add_rect(s, 0, Inches(5.55), SLIDE_W, Inches(1.95), RGBColor(0x0A, 0x30, 0x4C))
add_rect(s, 0, Inches(5.45), SLIDE_W, Inches(0.12), TEAL)
add_rect(s, Inches(0.0), 0, Inches(0.28), SLIDE_H, TEAL)

add_text(s, Inches(0.9), Inches(0.55), Inches(11.5), Inches(0.5),
         "KLINICZNE PODSTAWY FIZJOTERAPII W PULMONOLOGII",
         size=16, color=LIGHT_TEAL, bold=True)
add_text(s, Inches(0.9), Inches(1.75), Inches(11.5), Inches(2.3),
         "Zasady planowania i programowania\nfizjoterapii pacjentów z chorobami\nukładu oddechowego",
         size=34, color=WHITE, bold=True, line_spacing=1.05)

add_round_rect(s, Inches(0.9), Inches(4.35), Inches(5.6), Inches(0.85), TEAL)
add_text(s, Inches(1.05), Inches(4.35), Inches(5.4), Inches(0.85),
         "Jednostka chorobowa: ASTMA OSKRZELOWA",
         size=18, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)

add_text(s, Inches(0.9), Inches(5.75), Inches(11.5), Inches(0.5),
         "Metody badania pacjenta  •  Cele terapii  •  Plan i program fizjoterapii",
         size=15, color=LIGHT_TEAL)
add_text(s, Inches(0.9), Inches(6.45), Inches(11.5), Inches(0.5),
         "Praca projektowa  |  Fizjoterapia",
         size=13, color=RGBColor(0x9F, 0xC9, 0xD1))

# ===========================================================================
# SLAJD 2 - PLAN PREZENTACJI
# ===========================================================================
s = content_slide("Plan prezentacji", "Zakres omawianych zagadnień")
col1 = [
    "Definicja i epidemiologia astmy oskrzelowej",
    "Etiologia i czynniki ryzyka",
    "Patofizjologia i obraz kliniczny",
    "Klasyfikacja i przebieg choroby",
    "Diagnostyka lekarska – rozpoznanie",
    "Metody badania pacjenta w fizjoterapii",
]
col2 = [
    "Cele fizjoterapii w astmie",
    "Zasady planowania terapii",
    "Zasady programowania terapii",
    "Środki i metody fizjoterapeutyczne",
    "Przykładowy plan terapii",
    "Ocena efektów, wnioski, bibliografia",
]
add_round_rect(s, Inches(0.6), Inches(1.5), Inches(6.0), Inches(4.9), WHITE,
               line_color=RGBColor(0xD5, 0xE3, 0xE6))
add_round_rect(s, Inches(6.85), Inches(1.5), Inches(6.0), Inches(4.9), WHITE,
               line_color=RGBColor(0xD5, 0xE3, 0xE6))
add_text(s, Inches(0.9), Inches(1.7), Inches(5.4), Inches(0.5),
         "Część teoretyczna", size=17, color=NAVY, bold=True)
add_text(s, Inches(7.15), Inches(1.7), Inches(5.4), Inches(0.5),
         "Część fizjoterapeutyczna", size=17, color=NAVY, bold=True)
add_bullets(s, Inches(0.95), Inches(2.35), Inches(5.4), Inches(4.0),
            col1, size=16, space_after=12)
add_bullets(s, Inches(7.2), Inches(2.35), Inches(5.4), Inches(4.0),
            col2, size=16, space_after=12)

# ===========================================================================
# SLAJD 3 - DEFINICJA
# ===========================================================================
s = content_slide("Astma oskrzelowa – definicja", "Wg Global Initiative for Asthma (GINA)")
add_round_rect(s, Inches(0.6), Inches(1.55), Inches(12.1), Inches(1.85), LIGHT_TEAL)
add_text(s, Inches(0.95), Inches(1.75), Inches(11.4), Inches(1.5),
         "Astma to heterogenna choroba charakteryzująca się przewlekłym zapaleniem "
         "dróg oddechowych. Definiują ją występujące w wywiadzie objawy ze strony układu "
         "oddechowego – świszczący oddech, duszność, uczucie ściskania w klatce piersiowej "
         "i kaszel – o zmiennym nasileniu w czasie, wraz ze zmiennym ograniczeniem "
         "przepływu powietrza przez drogi oddechowe.",
         size=17, color=DARK_TEXT, italic=True, anchor=MSO_ANCHOR.MIDDLE,
         line_spacing=1.12)
add_text(s, Inches(0.95), Inches(3.5), Inches(11.4), Inches(0.4),
         "Cechy definiujące astmę:", size=17, color=NAVY, bold=True)
card(s, Inches(0.6), Inches(4.0), Inches(3.9), Inches(2.6),
     "Przewlekłe zapalenie", [
         "Zapalny charakter choroby",
         "Nacieki komórek zapalnych (eozynofile, limfocyty Th2, mastocyty)",
         "Nadreaktywność oskrzeli",
     ], body_size=13)
card(s, Inches(4.7), Inches(4.0), Inches(3.9), Inches(2.6),
     "Zmienność objawów", [
         "Nasilenie zmienne w czasie",
         "Objawy nocne i nad ranem",
         "Prowokacja przez czynniki wyzwalające",
     ], head_color=ACCENT, body_size=13)
card(s, Inches(8.8), Inches(4.0), Inches(3.9), Inches(2.6),
     "Obturacja odwracalna", [
         "Zmienne ograniczenie przepływu powietrza",
         "Odwracalność po lekach rozszerzających oskrzela",
         "Skurcz mięśni gładkich oskrzeli",
     ], head_color=NAVY, body_size=13)

# ===========================================================================
# SLAJD 4 - EPIDEMIOLOGIA
# ===========================================================================
s = content_slide("Epidemiologia", "Skala problemu zdrowotnego")
stats = [
    ("~262 mln", "osób na świecie choruje na astmę (dane WHO/GINA)"),
    ("~4 mln", "chorych na astmę w Polsce"),
    ("~455 tys.", "zgonów rocznie związanych z astmą na świecie"),
    ("Nr 1", "najczęstsza przewlekła choroba układu oddechowego u dzieci"),
]
x = Inches(0.6)
for val, desc in stats:
    add_round_rect(s, x, Inches(1.6), Inches(2.95), Inches(2.3), WHITE,
                   line_color=RGBColor(0xD5, 0xE3, 0xE6))
    add_rect(s, x, Inches(1.6), Inches(2.95), Inches(0.12), TEAL)
    add_text(s, x, Inches(1.95), Inches(2.95), Inches(0.9), val,
             size=32, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, x + Inches(0.15), Inches(2.85), Inches(2.65), Inches(1.0), desc,
             size=13, color=GREY_TEXT, align=PP_ALIGN.CENTER)
    x += Inches(3.05)

add_round_rect(s, Inches(0.6), Inches(4.25), Inches(12.1), Inches(2.25), LIGHT_TEAL)
add_text(s, Inches(0.9), Inches(4.4), Inches(11.5), Inches(0.4),
         "Znaczenie kliniczne i społeczne:", size=17, color=NAVY, bold=True)
add_bullets(s, Inches(0.95), Inches(4.9), Inches(11.4), Inches(1.5), [
    "Rosnąca częstość zachorowań, zwłaszcza w krajach uprzemysłowionych i w populacji miejskiej",
    "Istotne koszty bezpośrednie (leczenie) i pośrednie (absencja szkolna i zawodowa, hospitalizacje)",
    "Dobrze kontrolowana astma pozwala na normalne funkcjonowanie i aktywność fizyczną – kluczowa rola edukacji i rehabilitacji",
], size=15, space_after=7)

# ===========================================================================
# SLAJD 5 - ETIOLOGIA I CZYNNIKI RYZYKA
# ===========================================================================
s = content_slide("Etiologia i czynniki ryzyka")
card(s, Inches(0.6), Inches(1.5), Inches(3.9), Inches(5.0),
     "Czynniki osobnicze", [
         "Predyspozycja genetyczna",
         "Atopia (skłonność do IgE-zależnych reakcji)",
         "Płeć (u dzieci częściej chłopcy, u dorosłych kobiety)",
         "Otyłość",
         "Nadreaktywność oskrzeli",
     ], head_color=NAVY, body_size=14)
card(s, Inches(4.7), Inches(1.5), Inches(3.9), Inches(5.0),
     "Alergeny i środowisko", [
         "Roztocza kurzu domowego",
         "Pyłki roślin, zarodniki grzybów pleśniowych",
         "Sierść i naskórek zwierząt",
         "Dym tytoniowy (czynny i bierny)",
         "Zanieczyszczenie powietrza",
         "Alergeny zawodowe",
     ], head_color=TEAL, body_size=14)
card(s, Inches(8.8), Inches(1.5), Inches(3.9), Inches(5.0),
     "Czynniki wyzwalające", [
         "Infekcje układu oddechowego (wirusowe)",
         "Wysiłek fizyczny (astma wysiłkowa)",
         "Zimne powietrze",
         "Silne emocje, stres",
         "Leki (NLPZ, β-blokery)",
         "Refluks żołądkowo-przełykowy",
     ], head_color=ACCENT, body_size=14)

# ===========================================================================
# SLAJD 6 - PATOFIZJOLOGIA
# ===========================================================================
s = content_slide("Patofizjologia", "Mechanizmy obturacji oskrzeli")
add_bullets(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(4.8), [
    "Przewlekłe zapalenie ściany oskrzeli",
    ("Naciek eozynofilów, limfocytów Th2, mastocytów", 1),
    ("Uwalnianie mediatorów: histamina, leukotrieny, cytokiny", 1),
    "Skurcz mięśni gładkich oskrzeli (bronchospazm)",
    "Obrzęk błony śluzowej dróg oddechowych",
    "Nadmierne wydzielanie gęstego śluzu (czopy śluzowe)",
    "Nadreaktywność oskrzeli na bodźce",
    "Remodeling (przebudowa) ścian oskrzeli",
    ("Włóknienie podnabłonkowe, przerost mięśni – zmiany utrwalone", 1),
], size=16, space_after=8)
add_round_rect(s, Inches(6.9), Inches(1.6), Inches(5.8), Inches(4.8), LIGHT_TEAL)
add_text(s, Inches(7.15), Inches(1.8), Inches(5.3), Inches(0.5),
         "Skutki czynnościowe", size=17, color=NAVY, bold=True)
add_bullets(s, Inches(7.2), Inches(2.45), Inches(5.3), Inches(3.8), [
    "Zwężenie światła oskrzeli → wzrost oporu w drogach oddechowych",
    "Ograniczenie przepływu wydechowego (spadek FEV1, PEF)",
    "Rozdęcie płuc (hiperinflacja), pułapka powietrzna",
    "Wzrost pracy oddechowej i duszność",
    "Zaburzenia stosunku wentylacja/perfuzja → hipoksemia",
    "Odwracalność – kluczowa cecha różnicująca z POChP",
], size=15, space_after=9)

# ===========================================================================
# SLAJD 7 - OBRAZ KLINICZNY
# ===========================================================================
s = content_slide("Obraz kliniczny – objawy")
card(s, Inches(0.6), Inches(1.55), Inches(5.9), Inches(2.4),
     "Objawy podmiotowe (skargi chorego)", [
         "Napadowa duszność (głównie wydechowa)",
         "Świszczący oddech (wheezing)",
         "Uczucie ściskania / ucisku w klatce piersiowej",
         "Suchy, napadowy kaszel (często nocny)",
     ], head_color=TEAL, body_size=14)
card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.4),
     "Objawy przedmiotowe (badanie)", [
         "Świsty i furczenia w osłuchiwaniu",
         "Wydłużona faza wydechu",
         "Użycie dodatkowych mięśni oddechowych",
         "Tachykardia, w ciężkim napadzie sinica",
     ], head_color=NAVY, body_size=14)
add_round_rect(s, Inches(0.6), Inches(4.15), Inches(12.1), Inches(2.35), LIGHT_TEAL)
add_text(s, Inches(0.9), Inches(4.3), Inches(11.5), Inches(0.4),
         "Charakterystyka objawów:", size=16, color=NAVY, bold=True)
add_bullets(s, Inches(0.95), Inches(4.8), Inches(11.4), Inches(1.6), [
    "Zmienność w czasie – nasilenie zwłaszcza w nocy i nad ranem",
    "Występowanie / nasilenie po kontakcie z czynnikiem wyzwalającym (alergen, wysiłek, zimne powietrze)",
    "Ustępowanie samoistne lub po zastosowaniu leków rozszerzających oskrzela",
    "Zaostrzenie (napad astmy) – nagłe, znaczne nasilenie objawów, stan zagrożenia życia",
], size=14, space_after=6)

# ===========================================================================
# SLAJD 8 - KLASYFIKACJA
# ===========================================================================
s = content_slide("Klasyfikacja i stopnie kontroli astmy")
add_text(s, Inches(0.6), Inches(1.45), Inches(12), Inches(0.4),
         "Podział wg stopnia kontroli objawów (GINA):", size=17, color=NAVY, bold=True)
levels = [
    ("Astma\nkontrolowana", ACCENT, [
        "Objawy dzienne ≤ 2×/tydz.",
        "Bez ograniczeń aktywności",
        "Bez objawów nocnych",
        "Leki doraźne ≤ 2×/tydz.",
    ]),
    ("Astma częściowo\nkontrolowana", TEAL, [
        "Spełnione 1–2 kryteria",
        "z zakresu braku kontroli",
        "Okresowe objawy",
        "Wymaga modyfikacji leczenia",
    ]),
    ("Astma\nniekontrolowana", CORAL, [
        "Spełnione 3–4 kryteria",
        "Częste objawy dzienne i nocne",
        "Ograniczenie aktywności",
        "Częste stosowanie leków doraźnych",
    ]),
]
x = Inches(0.6)
for name, color, items in levels:
    add_round_rect(s, x, Inches(1.95), Inches(3.95), Inches(3.4), WHITE,
                   line_color=RGBColor(0xD5, 0xE3, 0xE6))
    add_round_rect(s, x, Inches(1.95), Inches(3.95), Inches(0.85), color)
    add_text(s, x, Inches(1.95), Inches(3.95), Inches(0.85), name,
             size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(s, x + Inches(0.25), Inches(3.0), Inches(3.5), Inches(2.2),
                items, size=13, space_after=6)
    x += Inches(4.05)
add_round_rect(s, Inches(0.6), Inches(5.6), Inches(12.1), Inches(0.95), LIGHT_TEAL)
add_text(s, Inches(0.9), Inches(5.6), Inches(11.5), Inches(0.95),
         "Astmę klasyfikuje się też wg ciężkości (leczenie potrzebne do uzyskania kontroli): "
         "lekka, umiarkowana, ciężka. Wyróżnia się m.in. astmę alergiczną, niealergiczną, "
         "wysiłkową oraz zawodową.",
         size=14, color=DARK_TEXT, anchor=MSO_ANCHOR.MIDDLE)

# ===========================================================================
# SLAJD 9 - DIAGNOSTYKA LEKARSKA
# ===========================================================================
s = content_slide("Diagnostyka – rozpoznanie lekarskie")
add_bullets(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(4.9), [
    "Wywiad – charakterystyczne, zmienne objawy",
    "Badanie przedmiotowe – osłuchiwanie",
    "Spirometria z próbą rozkurczową",
    ("Obturacja: FEV1/FVC poniżej normy", 1),
    ("Odwracalność: wzrost FEV1 ≥ 12% i ≥ 200 ml po leku rozkurczowym", 1),
    "Pomiar szczytowego przepływu wydechowego (PEF)",
    ("Zmienność dobowa PEF > 10%", 1),
    "Testy prowokacyjne (metacholina, wysiłek)",
    "Testy alergiczne (skórne, IgE swoiste)",
    "Ocena stanu zapalnego – FeNO, eozynofilia",
], size=15, space_after=6)
add_round_rect(s, Inches(6.9), Inches(1.6), Inches(5.8), Inches(4.9), LIGHT_TEAL)
add_text(s, Inches(7.15), Inches(1.8), Inches(5.3), Inches(0.5),
         "Spirometria – kluczowe parametry", size=16, color=NAVY, bold=True)
add_bullets(s, Inches(7.2), Inches(2.4), Inches(5.3), Inches(3.9), [
    "FEV1 – natężona objętość wydechowa pierwszosekundowa",
    "FVC – natężona pojemność życiowa",
    "FEV1/FVC – wskaźnik Tiffeneau (obturacja)",
    "PEF – szczytowy przepływ wydechowy (monitorowanie domowe)",
    "MEF – maksymalne przepływy wydechowe (małe oskrzela)",
], size=15, space_after=10)
add_text(s, Inches(7.2), Inches(5.85), Inches(5.3), Inches(0.6),
         "Rozpoznanie stawia lekarz – fizjoterapeuta wykorzystuje wyniki w planowaniu terapii.",
         size=12, color=GREY_TEXT, italic=True)

# ===========================================================================
# SLAJD 10 - METODY BADANIA PACJENTA (wprowadzenie / wywiad)
# ===========================================================================
s = content_slide("Metody badania pacjenta w fizjoterapii", "Badanie podmiotowe (wywiad)")
add_bullets(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(4.9), [
    "Wywiad chorobowy i fizjoterapeutyczny:",
    ("Czas trwania i przebieg choroby", 1),
    ("Częstość i okoliczności napadów duszności", 1),
    ("Czynniki wyzwalające objawy", 1),
    ("Tolerancja wysiłku fizycznego", 1),
    ("Objawy nocne i poranne", 1),
    ("Stosowane leki (wziewne, doraźne)", 1),
    ("Choroby współistniejące (np. alergia, otyłość)", 1),
    ("Nałogi – palenie tytoniu", 1),
    ("Aktywność fizyczna i tryb życia", 1),
    ("Oczekiwania i cele pacjenta", 1),
], size=15, space_after=5)
add_round_rect(s, Inches(6.9), Inches(1.6), Inches(5.8), Inches(4.9), LIGHT_TEAL)
add_text(s, Inches(7.15), Inches(1.8), Inches(5.3), Inches(0.5),
         "Narzędzia oceny subiektywnej", size=16, color=NAVY, bold=True)
add_bullets(s, Inches(7.2), Inches(2.4), Inches(5.3), Inches(3.9), [
    "ACT – Test Kontroli Astmy (Asthma Control Test)",
    "ACQ – Asthma Control Questionnaire",
    "AQLQ – kwestionariusz jakości życia w astmie",
    "Skala duszności mMRC (Modified Medical Research Council)",
    "Skala Borga – ocena duszności i zmęczenia wysiłkowego",
    "Dzienniczek objawów i pomiarów PEF",
], size=15, space_after=9)

# ===========================================================================
# SLAJD 11 - BADANIE PRZEDMIOTOWE / FIZYKALNE
# ===========================================================================
s = content_slide("Metody badania pacjenta – badanie przedmiotowe")
card(s, Inches(0.6), Inches(1.55), Inches(5.9), Inches(2.55),
     "Ocena wzrokowa i palpacyjna", [
         "Tor i częstość oddychania (norma 12–16/min)",
         "Typ oddychania (piersiowy / brzuszny / mieszany)",
         "Symetria ruchów klatki piersiowej",
         "Użycie pomocniczych mięśni oddechowych",
         "Postawa ciała, ustawienie obręczy barkowej",
     ], head_color=TEAL, body_size=13)
card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(2.55),
     "Pomiary i osłuchiwanie", [
         "Ruchomość oddechowa klatki piersiowej (cyrtometria)",
         "Osłuchiwanie – świsty, furczenia, wydłużony wydech",
         "Saturacja krwi tętniczej (pulsoksymetria SpO₂)",
         "Częstość tętna, ciśnienie tętnicze",
         "Ocena siły mięśni oddechowych (PImax, PEmax)",
     ], head_color=NAVY, body_size=13)
add_round_rect(s, Inches(0.6), Inches(4.3), Inches(12.1), Inches(2.2), LIGHT_TEAL)
add_text(s, Inches(0.9), Inches(4.45), Inches(11.5), Inches(0.4),
         "Ocena tolerancji wysiłku i wydolności:", size=16, color=NAVY, bold=True)
add_bullets(s, Inches(0.95), Inches(4.95), Inches(11.4), Inches(1.5), [
    "6-minutowy test marszu (6MWT) – dystans, duszność, saturacja przed i po wysiłku",
    "Test wahadłowy (shuttle walk test), próba wysiłkowa na ergometrze",
    "Ocena postawy ciała i napięcia mięśniowego (mięśnie pomocnicze, przepona)",
], size=14, space_after=7)

# ===========================================================================
# SLAJD 12 - CELE FIZJOTERAPII
# ===========================================================================
s = content_slide("Cele fizjoterapii w astmie oskrzelowej")
add_text(s, Inches(0.6), Inches(1.4), Inches(12), Inches(0.4),
         "Fizjoterapia stanowi element kompleksowej rehabilitacji pulmonologicznej – "
         "uzupełnia (nie zastępuje) farmakoterapię.", size=15, color=GREY_TEXT, italic=True)
goals = [
    ("Cele główne", TEAL, [
        "Poprawa kontroli objawów astmy",
        "Zwiększenie tolerancji wysiłku i wydolności",
        "Poprawa jakości życia",
        "Zmniejszenie liczby zaostrzeń",
    ]),
    ("Cele oddechowe", NAVY, [
        "Nauka prawidłowego toru oddychania",
        "Zwiększenie ruchomości klatki piersiowej",
        "Wzmocnienie mięśni oddechowych",
        "Ułatwienie ewakuacji wydzieliny",
    ]),
    ("Cele funkcjonalne", ACCENT, [
        "Korekcja postawy ciała",
        "Redukcja duszności i lęku",
        "Zwiększenie aktywności fizycznej",
        "Edukacja i samodzielność chorego",
    ]),
]
x = Inches(0.6)
for name, color, items in goals:
    add_round_rect(s, x, Inches(2.0), Inches(3.95), Inches(4.4), WHITE,
                   line_color=RGBColor(0xD5, 0xE3, 0xE6))
    add_round_rect(s, x, Inches(2.0), Inches(3.95), Inches(0.6), color)
    add_text(s, x, Inches(2.0), Inches(3.95), Inches(0.6), name,
             size=16, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(s, x + Inches(0.25), Inches(2.85), Inches(3.5), Inches(3.4),
                items, size=14, space_after=11)
    x += Inches(4.05)

# ===========================================================================
# SLAJD 13 - ZASADY PLANOWANIA
# ===========================================================================
s = content_slide("Zasady planowania fizjoterapii")
add_bullets(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(4.9), [
    "Indywidualizacja – dostosowanie do pacjenta",
    ("Wiek, stopień ciężkości i kontroli astmy", 1),
    ("Choroby współistniejące, wydolność", 1),
    "Kompleksowość – współpraca zespołu",
    ("Lekarz, fizjoterapeuta, pielęgniarka, psycholog", 1),
    "Etapowość i stopniowanie obciążeń",
    "Systematyczność i regularność ćwiczeń",
    "Ciągłość procesu usprawniania",
    "Bezpieczeństwo – kontrola objawów i saturacji",
    "Ocena wyjściowa i cele SMART",
], size=16, space_after=6)
add_round_rect(s, Inches(6.9), Inches(1.6), Inches(5.8), Inches(4.9), LIGHT_TEAL)
add_text(s, Inches(7.15), Inches(1.8), Inches(5.3), Inches(0.5),
         "Etapy planowania terapii", size=16, color=NAVY, bold=True)
steps = [
    "1. Ocena stanu pacjenta (badanie podmiotowe i przedmiotowe)",
    "2. Identyfikacja problemów i deficytów funkcjonalnych",
    "3. Ustalenie celów krótko- i długoterminowych",
    "4. Dobór metod i środków fizjoterapii",
    "5. Realizacja programu terapii",
    "6. Ocena efektów i modyfikacja planu",
]
add_bullets(s, Inches(7.2), Inches(2.45), Inches(5.3), Inches(3.9),
            steps, size=15, space_after=11, bullet_color=NAVY)

# ===========================================================================
# SLAJD 14 - ZASADY PROGRAMOWANIA
# ===========================================================================
s = content_slide("Zasady programowania fizjoterapii", "Parametry dawkowania (model FITT)")
fitt = [
    ("F – Frequency\n(częstotliwość)", TEAL,
     "Trening wytrzymałościowy 3–5×/tydz.; ćwiczenia oddechowe codziennie; "
     "trening oporowy 2–3×/tydz."),
    ("I – Intensity\n(intensywność)", NAVY,
     "Umiarkowana: 40–70% rezerwy tętna; duszność 3–4 w skali Borga (0–10); "
     "stopniowe zwiększanie obciążeń"),
    ("T – Time\n(czas)", ACCENT,
     "20–60 min sesji aerobowej; program rehabilitacji 6–12 tygodni; "
     "ćwiczenia oddechowe kilka razy dziennie po kilka minut"),
    ("T – Type\n(rodzaj)", CORAL,
     "Trening aerobowy (marsz, rower, pływanie), oporowy, ćwiczenia oddechowe, "
     "trening mięśni wdechowych"),
]
positions = [(Inches(0.6), Inches(1.6)), (Inches(6.75), Inches(1.6)),
             (Inches(0.6), Inches(4.15)), (Inches(6.75), Inches(4.15))]
for (px, py), (name, color, desc) in zip(positions, fitt):
    add_round_rect(s, px, py, Inches(6.0), Inches(2.3), WHITE,
                   line_color=RGBColor(0xD5, 0xE3, 0xE6))
    add_round_rect(s, px, py, Inches(1.9), Inches(2.3), color)
    add_text(s, px, py, Inches(1.9), Inches(2.3), name, size=16, color=WHITE,
             bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, px + Inches(2.1), py + Inches(0.15), Inches(3.75), Inches(2.0),
             desc, size=14, color=DARK_TEXT, anchor=MSO_ANCHOR.MIDDLE,
             line_spacing=1.1)

# ===========================================================================
# SLAJD 15 - PRZEGLAD SRODKOW FIZJOTERAPII
# ===========================================================================
s = content_slide("Środki i metody fizjoterapii – przegląd")
groups = [
    ("Kinezyterapia oddechowa", [
        "Ćwiczenia oddechowe",
        "Nauka toru oddychania",
        "Trening mięśni oddechowych",
        "Ćwiczenia rozprężające",
    ], TEAL),
    ("Techniki oczyszczania", [
        "Drenaż ułożeniowy",
        "Techniki natężonego wydechu",
        "Urządzenia PEP / oscylacyjne",
        "Aktywny cykl oddechowy (ACBT)",
    ], NAVY),
    ("Trening fizyczny", [
        "Trening aerobowy",
        "Trening oporowy",
        "Ćwiczenia ogólnousprawniające",
        "Rehabilitacja pulmonologiczna",
    ], ACCENT),
    ("Metody wspomagające", [
        "Edukacja pacjenta",
        "Techniki relaksacyjne",
        "Korekcja postawy ciała",
        "Fizykoterapia (inhalacje)",
    ], CORAL),
]
x = Inches(0.6)
for name, items, color in groups:
    add_round_rect(s, x, Inches(1.55), Inches(2.95), Inches(4.9), WHITE,
                   line_color=RGBColor(0xD5, 0xE3, 0xE6))
    add_round_rect(s, x, Inches(1.55), Inches(2.95), Inches(0.95), color)
    add_text(s, x + Inches(0.1), Inches(1.55), Inches(2.75), Inches(0.95), name,
             size=14, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(s, x + Inches(0.2), Inches(2.7), Inches(2.6), Inches(3.6),
                items, size=13, space_after=10)
    x += Inches(3.05)

# ===========================================================================
# SLAJD 16 - CWICZENIA ODDECHOWE
# ===========================================================================
s = content_slide("Ćwiczenia oddechowe (breathing retraining)")
add_bullets(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(4.9), [
    "Oddychanie przeponowe (torem brzusznym)",
    ("Angażowanie przepony, redukcja pracy oddechowej", 1),
    "Oddychanie przez „zasznurowane usta” (pursed-lip breathing)",
    ("Wydłużenie wydechu, przeciwdziałanie zapadaniu oskrzeli", 1),
    "Kontrola tempa i rytmu oddechu (spowolnienie)",
    "Metoda Butejki – redukcja hiperwentylacji",
    "Ćwiczenia wg techniki Papworth",
    "Elementy jogi i ćwiczeń relaksacyjnych oddechu",
    "Ćwiczenia rozprężające dolne partie płuc",
], size=15, space_after=6)
add_round_rect(s, Inches(6.9), Inches(1.6), Inches(5.8), Inches(4.9), LIGHT_TEAL)
add_text(s, Inches(7.15), Inches(1.8), Inches(5.3), Inches(0.5),
         "Efekty i dowody naukowe", size=16, color=NAVY, bold=True)
add_bullets(s, Inches(7.2), Inches(2.4), Inches(5.3), Inches(3.9), [
    "Poprawa kontroli objawów i jakości życia (przegląd Cochrane, Santino i wsp. 2020)",
    "Redukcja objawów lęku i hiperwentylacji",
    "Zmniejszenie zużycia leków doraźnych",
    "Poprawa wzorca oddechowego i tolerancji wysiłku",
    "Techniki oddechowe zalecane jako uzupełnienie farmakoterapii (GINA)",
], size=15, space_after=10)

# ===========================================================================
# SLAJD 17 - TRENING MIESNI ODDECHOWYCH + OCZYSZCZANIE
# ===========================================================================
s = content_slide("Trening mięśni oddechowych i oczyszczanie oskrzeli")
card(s, Inches(0.6), Inches(1.55), Inches(5.9), Inches(4.9),
     "Trening mięśni oddechowych (IMT)", [
         "Trening mięśni wdechowych z użyciem oporu (np. Threshold IMT)",
         "Wzmacnia przeponę i mięśnie międzyżebrowe",
         "Typowo: obciążenie 30–50% PImax",
         "Zmniejsza duszność i pracę oddechową",
         "Poprawa siły i wytrzymałości mięśni oddechowych",
         "Ćwiczenia oporowe mięśni wydechowych w razie wskazań",
     ], head_color=TEAL, body_size=14)
card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(4.9),
     "Techniki oczyszczania drzewa oskrzelowego", [
         "Wskazane głównie przy nadmiernej wydzielinie / zaostrzeniu",
         "Aktywny cykl technik oddechowych (ACBT)",
         "Technika natężonego wydechu (huffing, ‘forced expiration’)",
         "Drenaż ułożeniowy wspomagany oklepywaniem/wibracjami",
         "Urządzenia PEP i oscylacyjne (flutter, acapella)",
         "Efektywny kaszel – nauka kontrolowanego odkrztuszania",
     ], head_color=NAVY, body_size=14)

# ===========================================================================
# SLAJD 18 - TRENING FIZYCZNY / REHABILITACJA
# ===========================================================================
s = content_slide("Trening fizyczny i rehabilitacja pulmonologiczna")
add_bullets(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(4.9), [
    "Trening aerobowy (wytrzymałościowy):",
    ("Marsz, nordic walking, rower, pływanie", 1),
    ("Poprawa wydolności krążeniowo-oddechowej", 1),
    "Trening oporowy (siłowy):",
    ("Wzmacnianie dużych grup mięśniowych", 1),
    ("Przeciwdziałanie skutkom steroidoterapii", 1),
    "Ćwiczenia ogólnousprawniające i rozciągające",
    "Rozgrzewka i wychłodzenie – ważne w astmie wysiłkowej",
    "Stopniowa progresja obciążeń, monitoring objawów",
], size=15, space_after=6)
add_round_rect(s, Inches(6.9), Inches(1.6), Inches(5.8), Inches(4.9), LIGHT_TEAL)
add_text(s, Inches(7.15), Inches(1.8), Inches(5.3), Inches(0.5),
         "Korzyści z treningu fizycznego", size=16, color=NAVY, bold=True)
add_bullets(s, Inches(7.2), Inches(2.4), Inches(5.3), Inches(3.9), [
    "Wzrost wydolności fizycznej i tolerancji wysiłku (Cochrane, Carson i wsp. 2013)",
    "Zmniejszenie duszności wysiłkowej",
    "Poprawa jakości życia i samopoczucia",
    "Redukcja objawów astmy i częstości zaostrzeń",
    "Aktywność fizyczna jest bezpieczna w dobrze kontrolowanej astmie",
], size=15, space_after=10)

# ===========================================================================
# SLAJD 19 - ASTMA WYSILKOWA
# ===========================================================================
s = content_slide("Astma wysiłkowa – szczególne zasady", "Exercise-Induced Bronchoconstriction")
add_round_rect(s, Inches(0.6), Inches(1.55), Inches(12.1), Inches(1.3), LIGHT_TEAL)
add_text(s, Inches(0.95), Inches(1.55), Inches(11.4), Inches(1.3),
         "Skurcz oskrzeli wywołany wysiłkiem (EIB) – przemijające zwężenie dróg oddechowych "
         "pojawiające się w trakcie lub po intensywnym wysiłku, nasilane przez zimne, suche powietrze. "
         "Nie jest przeciwwskazaniem do aktywności – wymaga odpowiedniego postępowania.",
         size=15, color=DARK_TEXT, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)
card(s, Inches(0.6), Inches(3.05), Inches(5.9), Inches(3.4),
     "Postępowanie profilaktyczne", [
         "Stosowanie leku rozkurczowego przed wysiłkiem (wg zaleceń lekarza)",
         "Dokładna, stopniowana rozgrzewka (‘refractory period’)",
         "Powolne wychładzanie po wysiłku",
         "Unikanie wysiłku w zimnym, suchym i zanieczyszczonym powietrzu",
         "Oddychanie przez nos, ochrona dróg oddechowych (np. komin/maska zimą)",
     ], head_color=TEAL, body_size=14)
card(s, Inches(6.8), Inches(3.05), Inches(5.9), Inches(3.4),
     "Zalecane formy aktywności", [
         "Pływanie (ciepłe, wilgotne powietrze) – dobrze tolerowane",
         "Aktywności o zmiennej intensywności (gry zespołowe)",
         "Marsz, jazda na rowerze w umiarkowanym tempie",
         "Sporty wymagające krótkich zrywów lepiej tolerowane niż długi wysiłek ciągły w zimnie",
     ], head_color=NAVY, body_size=14)

# ===========================================================================
# SLAJD 20 - EDUKACJA PACJENTA
# ===========================================================================
s = content_slide("Edukacja pacjenta i profilaktyka")
add_bullets(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(4.9), [
    "Wiedza o chorobie i mechanizmie objawów",
    "Rozpoznawanie i unikanie czynników wyzwalających",
    "Prawidłowa technika inhalacji leków",
    "Samokontrola – pomiar PEF, dzienniczek objawów",
    "Rozpoznawanie objawów zaostrzenia (plan działania)",
    "Znaczenie systematycznej aktywności fizycznej",
    "Techniki radzenia sobie z dusznością i lękiem",
    "Kontrola masy ciała, zaprzestanie palenia",
], size=16, space_after=8)
add_round_rect(s, Inches(6.9), Inches(1.6), Inches(5.8), Inches(4.9), LIGHT_TEAL)
add_text(s, Inches(7.15), Inches(1.8), Inches(5.3), Inches(0.5),
         "Pisemny plan działania w astmie", size=16, color=NAVY, bold=True)
add_bullets(s, Inches(7.2), Inches(2.45), Inches(5.3), Inches(3.9), [
    "Strefa zielona – astma kontrolowana, kontynuacja leczenia",
    "Strefa żółta – narastanie objawów, modyfikacja leków wg zaleceń",
    "Strefa czerwona – zaostrzenie, pilny kontakt z lekarzem / pomoc doraźna",
    "Edukacja zwiększa współpracę (compliance) i skuteczność terapii",
], size=15, space_after=12, bullet_color=NAVY)

# ===========================================================================
# SLAJD 21 - FIZJOTERAPIA W ZAOSTRZENIU vs STABILNA
# ===========================================================================
s = content_slide("Fizjoterapia – zaostrzenie vs okres stabilny")
card(s, Inches(0.6), Inches(1.55), Inches(5.9), Inches(4.9),
     "Okres zaostrzenia (napad)", [
         "Priorytet: farmakoterapia i tlenoterapia",
         "Pozycje ułatwiające oddychanie (np. woźnicy, wysoka)",
         "Oddychanie przez zasznurowane usta, spowolnienie oddechu",
         "Techniki relaksacyjne, redukcja lęku",
         "Delikatna ewakuacja wydzieliny (jeśli obecna)",
         "UNIKAĆ intensywnego wysiłku i forsownych technik",
     ], head_color=CORAL, body_size=14)
card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(4.9),
     "Okres stabilny (kontrola astmy)", [
         "Trening aerobowy i oporowy (progresja obciążeń)",
         "Ćwiczenia oddechowe i trening mięśni oddechowych",
         "Poprawa ruchomości klatki piersiowej i postawy",
         "Edukacja i utrwalanie nawyków ruchowych",
         "Techniki oczyszczania w razie zalegania wydzieliny",
         "Regularna aktywność fizyczna jako profilaktyka",
     ], head_color=ACCENT, body_size=14)

# ===========================================================================
# SLAJD 22 - PRZYKLADOWY PLAN TERAPII (opis pacjenta + cele)
# ===========================================================================
s = content_slide("Przykładowy plan terapii (1/2)", "Opis przypadku i cele")
add_round_rect(s, Inches(0.6), Inches(1.55), Inches(12.1), Inches(1.85), LIGHT_TEAL)
add_text(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(0.4),
         "Charakterystyka pacjenta:", size=16, color=NAVY, bold=True)
add_text(s, Inches(0.95), Inches(2.2), Inches(11.4), Inches(1.1),
         "Kobieta, 34 lata, astma oskrzelowa alergiczna, częściowo kontrolowana. "
         "Zgłasza duszność wysiłkową, obniżoną tolerancję wysiłku, epizody kaszlu nocnego. "
         "Prowadzi siedzący tryb życia. Bez zaostrzenia w chwili badania (astma stabilna).",
         size=15, color=DARK_TEXT, line_spacing=1.15)
card(s, Inches(0.6), Inches(3.6), Inches(5.9), Inches(2.9),
     "Cele krótkoterminowe (4–6 tyg.)", [
         "Nauka prawidłowego toru oddychania",
         "Opanowanie technik kontroli duszności",
         "Zwiększenie ruchomości klatki piersiowej",
         "Wprowadzenie regularnej aktywności aerobowej",
     ], head_color=TEAL, body_size=14)
card(s, Inches(6.8), Inches(3.6), Inches(5.9), Inches(2.9),
     "Cele długoterminowe (3–6 mies.)", [
         "Poprawa tolerancji wysiłku (dystans 6MWT)",
         "Lepsza kontrola astmy (wzrost wyniku ACT)",
         "Poprawa jakości życia (AQLQ)",
         "Utrwalenie nawyku aktywności fizycznej",
     ], head_color=NAVY, body_size=14)

# ===========================================================================
# SLAJD 23 - PRZYKLADOWY PLAN TERAPII - PROGRAM
# ===========================================================================
s = content_slide("Przykładowy plan terapii (2/2)", "Program tygodniowy")
add_text(s, Inches(0.6), Inches(1.4), Inches(12), Inches(0.4),
         "Program rehabilitacji – 8 tygodni, 3 sesje nadzorowane + zalecenia domowe:",
         size=15, color=GREY_TEXT, italic=True)
plan = [
    ("Ćwiczenia oddechowe", "Codziennie", "Przeponowe + pursed-lip, 2–3× dziennie po 5–10 min"),
    ("Trening mięśni wdechowych (IMT)", "5–6×/tydz.", "Threshold IMT, 30% PImax, 2 serie × 15 oddechów"),
    ("Trening aerobowy", "3–5×/tydz.", "Marsz/rower, 20–40 min, 40–70% rezerwy tętna, Borg 3–4"),
    ("Trening oporowy", "2–3×/tydz.", "Duże grupy mięśni, 1–3 serie × 8–12 powtórzeń"),
    ("Ćwiczenia klatki i postawy", "Codziennie", "Rozprężające, rozciągające, korekcja postawy"),
    ("Edukacja i samokontrola", "Ciągle", "Technika inhalacji, dzienniczek PEF, plan działania"),
]
# tabela
tx, ty = Inches(0.6), Inches(2.0)
col_w = [Inches(4.0), Inches(2.3), Inches(5.8)]
row_h = Inches(0.68)
headers = ["Element programu", "Częstotliwość", "Dawkowanie / uwagi"]
cx = tx
for i, htxt in enumerate(headers):
    add_rect(s, cx, ty, col_w[i], row_h, NAVY)
    add_text(s, cx + Inches(0.1), ty, col_w[i] - Inches(0.2), row_h, htxt,
             size=14, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    cx += col_w[i]
for r, (a, b, c) in enumerate(plan):
    ry = ty + row_h * (r + 1)
    bg = WHITE if r % 2 == 0 else LIGHT_TEAL
    cx = tx
    for i, val in enumerate((a, b, c)):
        add_rect(s, cx, ry, col_w[i], row_h, bg,
                 line_color=RGBColor(0xCF, 0xDD, 0xE0))
        add_text(s, cx + Inches(0.1), ry, col_w[i] - Inches(0.2), row_h, val,
                 size=12, color=DARK_TEXT, bold=(i == 0),
                 anchor=MSO_ANCHOR.MIDDLE)
        cx += col_w[i]

# ===========================================================================
# SLAJD 24 - PRZECIWWSKAZANIA I BEZPIECZENSTWO
# ===========================================================================
s = content_slide("Przeciwwskazania i zasady bezpieczeństwa")
card(s, Inches(0.6), Inches(1.55), Inches(5.9), Inches(4.9),
     "Przeciwwskazania / ostrożność", [
         "Ciężkie, niekontrolowane zaostrzenie astmy",
         "Znaczna hipoksemia (niska saturacja) bez zabezpieczenia",
         "Niestabilne choroby sercowo-naczyniowe",
         "Ostra infekcja z gorączką",
         "Wysiłek w środowisku wyzwalającym objawy (zimno, alergeny)",
         "Brak dostępu do leku doraźnego podczas ćwiczeń",
     ], head_color=CORAL, body_size=14)
card(s, Inches(6.8), Inches(1.55), Inches(5.9), Inches(4.9),
     "Zasady bezpieczeństwa", [
         "Ćwiczenia przy dobrej kontroli astmy",
         "Lek rozkurczowy dostępny podczas terapii",
         "Monitorowanie SpO₂, tętna i duszności (Borg)",
         "Rozgrzewka i stopniowa progresja obciążeń",
         "Przerwanie ćwiczeń przy nasileniu objawów",
         "Dostosowanie środowiska (temperatura, wilgotność, brak alergenów)",
     ], head_color=ACCENT, body_size=14)

# ===========================================================================
# SLAJD 25 - OCENA EFEKTOW TERAPII
# ===========================================================================
s = content_slide("Ocena efektów fizjoterapii")
add_text(s, Inches(0.6), Inches(1.4), Inches(12), Inches(0.4),
         "Monitorowanie postępów przy pomocy obiektywnych i subiektywnych mierników:",
         size=15, color=GREY_TEXT, italic=True)
groups = [
    ("Parametry czynnościowe", [
        "Spirometria (FEV1, PEF)",
        "Zmienność dobowa PEF",
        "Saturacja SpO₂",
        "Siła mięśni oddechowych (PImax)",
    ], TEAL),
    ("Wydolność i sprawność", [
        "6-minutowy test marszu (6MWT)",
        "Tolerancja wysiłku",
        "Duszność wysiłkowa (skala Borga)",
        "Ruchomość klatki piersiowej",
    ], NAVY),
    ("Kontrola i jakość życia", [
        "Test kontroli astmy (ACT / ACQ)",
        "Kwestionariusz AQLQ",
        "Częstość objawów i zaostrzeń",
        "Zużycie leków doraźnych",
    ], ACCENT),
]
x = Inches(0.6)
for name, items, color in groups:
    add_round_rect(s, x, Inches(2.0), Inches(3.95), Inches(4.4), WHITE,
                   line_color=RGBColor(0xD5, 0xE3, 0xE6))
    add_round_rect(s, x, Inches(2.0), Inches(3.95), Inches(0.75), color)
    add_text(s, x + Inches(0.1), Inches(2.0), Inches(3.75), Inches(0.75), name,
             size=15, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    add_bullets(s, x + Inches(0.25), Inches(3.0), Inches(3.5), Inches(3.3),
                items, size=14, space_after=12)
    x += Inches(4.05)

# ===========================================================================
# SLAJD 26 - PODSUMOWANIE / WNIOSKI
# ===========================================================================
s = content_slide("Podsumowanie i wnioski")
add_bullets(s, Inches(0.7), Inches(1.55), Inches(12.0), Inches(5.2), [
    "Astma oskrzelowa to przewlekła, zapalna choroba dróg oddechowych o zmiennym przebiegu i odwracalnej obturacji.",
    "Rzetelne badanie pacjenta (wywiad, badanie przedmiotowe, testy czynnościowe i skale) jest podstawą planowania terapii.",
    "Fizjoterapia jest ważnym, uzupełniającym elementem kompleksowego leczenia astmy – nie zastępuje farmakoterapii.",
    "Kluczowe środki: ćwiczenia oddechowe, trening mięśni oddechowych, trening fizyczny, techniki oczyszczania i edukacja.",
    "Planowanie i programowanie musi być indywidualne, etapowe, systematyczne i bezpieczne (model FITT, cele SMART).",
    "Regularna, właściwie dozowana aktywność fizyczna poprawia wydolność, kontrolę objawów i jakość życia chorych.",
    "Efekty należy monitorować obiektywnymi i subiektywnymi narzędziami oraz modyfikować program w oparciu o wyniki.",
], size=16, space_after=11)

# ===========================================================================
# SLAJD 27 - BIBLIOGRAFIA
# ===========================================================================
s = content_slide("Bibliografia", "Piśmiennictwo naukowe (PubMed / Google Scholar) i podręczniki")
refs_left = [
    "1. Global Initiative for Asthma (GINA). Global Strategy for Asthma Management and Prevention. 2023. www.ginasthma.org",
    "2. Santino TA, Chaves GSS, Freitas DA, Fregonezi GAF, Mendonça KMPP. Breathing exercises for adults with asthma. "
    "Cochrane Database Syst Rev. 2020;3:CD001277.",
    "3. Bruton A, Lee A, Yardley L, et al. Physiotherapy breathing retraining for asthma: a randomised controlled trial. "
    "Lancet Respir Med. 2018;6(1):19–28.",
    "4. Carson KV, Chandratilleke MG, Picot J, et al. Physical training for asthma. Cochrane Database Syst Rev. "
    "2013;(9):CD001116.",
    "5. Freitas DA, Holloway EA, Bruno SS, et al. Breathing exercises for adults with asthma. "
    "Cochrane Database Syst Rev. 2013;(10):CD001277.",
]
refs_right = [
    "6. Thomas M, Bruton A. Breathing exercises for asthma. Breathe. 2014;10(4):312–322.",
    "7. Hough A. Physiotherapy in Respiratory and Cardiac Care: An Evidence-Based Approach. "
    "Cengage Learning, 2014.",
    "8. Woźniewski M. (red.). Fizjoterapia w chorobach wewnętrznych. Wyd. Lekarskie PZWL, Warszawa 2021.",
    "9. Rosławski A., Woźniewski M. Fizjoterapia oddechowa. AWF Wrocław.",
    "10. Kwolek A. (red.). Rehabilitacja medyczna. Edra Urban & Partner, Wrocław.",
    "11. Bott J. i wsp. Guidelines for the physiotherapy management of the adult, medical, spontaneously breathing "
    "patient. Thorax. 2009;64(Suppl 1):i1–i52.",
]
add_round_rect(s, Inches(0.5), Inches(1.45), Inches(6.15), Inches(5.15), WHITE,
               line_color=RGBColor(0xD5, 0xE3, 0xE6))
add_round_rect(s, Inches(6.8), Inches(1.45), Inches(6.05), Inches(5.15), WHITE,
               line_color=RGBColor(0xD5, 0xE3, 0xE6))
add_bullets(s, Inches(0.75), Inches(1.65), Inches(5.7), Inches(4.8),
            refs_left, size=12.5, space_after=11, bullet_color=NAVY,
            line_spacing=1.05)
add_bullets(s, Inches(7.05), Inches(1.65), Inches(5.6), Inches(4.8),
            refs_right, size=12.5, space_after=11, bullet_color=NAVY,
            line_spacing=1.05)

# ===========================================================================
# SLAJD 28 - SLAJD KONCOWY
# ===========================================================================
s = add_slide()
set_bg(s, NAVY)
add_rect(s, 0, Inches(3.05), SLIDE_W, Inches(0.1), TEAL)
add_text(s, Inches(1.0), Inches(2.15), Inches(11.3), Inches(1.0),
         "Dziękuję za uwagę", size=44, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER)
add_text(s, Inches(1.0), Inches(3.35), Inches(11.3), Inches(0.6),
         "Zasady planowania i programowania fizjoterapii – Astma oskrzelowa",
         size=18, color=LIGHT_TEAL, align=PP_ALIGN.CENTER)
add_text(s, Inches(1.0), Inches(4.15), Inches(11.3), Inches(0.6),
         "Kliniczne podstawy fizjoterapii w pulmonologii",
         size=15, color=RGBColor(0x9F, 0xC9, 0xD1), align=PP_ALIGN.CENTER)

# ---------------------------------------------------------------------------
out = "Astma_oskrzelowa_fizjoterapia_prezentacja.pptx"
prs.save(out)
print("Zapisano:", out)
print("Liczba slajdów:", len(prs.slides._sldIdLst))
